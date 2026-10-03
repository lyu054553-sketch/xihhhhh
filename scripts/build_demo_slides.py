"""Build v1.2 HTML and editable PowerPoint from one source. Requires python-pptx."""

from html import escape
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SLIDES = [
    {
        "eyebrow": "HUO BU YA QIAN / LIVE BACKEND v1.2",
        "title": "货不压钱\n让库存决策有据可查",
        "subtitle": "连锁零售库存工作台 · 实际后端计算 · 合成零食数据",
        "points": [
            "先核对事实、数量与业务约束",
            "保存方案，人工审批后生成待执行任务",
            "采购情景先确认参数，再比较结果"
        ],
        "note": "当前为确定性规则工具 · 没有真实模型或外部 ERP 调用",
        "speaker": "本轮以魏的 lanyangyang 分支和 v1.2 当前契约联调。页面访问实际 FastAPI 和数据库，不以响应夹具生成成功结果。数据仍然是合成的，演示不代表真实客户收益。"
    },
    {
        "eyebrow": "01 / DAILY WORK",
        "title": "从需要关注的商品\n走到一份可核对方案",
        "subtitle": "经营总览 → 风险证据 → 工作台计算 → 人工审批 → 待执行任务",
        "points": [
            "总览分开显示账户、库存成本与采购付款",
            "风险详情保留事实、原因假设与缺项",
            "调拨、近效期、采购各自核对约束",
            "今日待办区分待审批、跟进与人工完成"
        ],
        "note": "调拨改变库存位置，不直接产生现金到账",
        "speaker": "用西湖文三店每日坚果礼盒作为演示。先核对证据，再进入调拨工作台修改数量。输入变化以后必须重新计算，不能继续使用旧结果。"
    },
    {
        "eyebrow": "02 / COMPUTATION",
        "title": "程序计算数量与金额\n前端呈现证据与错误",
        "subtitle": "当前使用魏的实际零售、工作台、方案和执行任务 API。",
        "points": [
            "金额为人民币元，百分比为 0—100",
            "初值和结果从后端返回，前端不补算",
            "无效方案展示约束错误，不能保存",
            "未知保留 null，真实缺项不补合成数字"
        ],
        "note": "当前依据根 API_CONTRACT.md 与 FastAPI 1.2.0",
        "speaker": "业务结果由确定性程序计算。三类工作台复用保存、提交、审批和任务流程。旧统一 Agent Run 草案不再是当前协议，也不会通过前端兼容层伪装成后端已经实现。"
    },
    {
        "eyebrow": "03 / HUMAN REVIEW",
        "title": "保存、审批、执行\n是三个不同的状态",
        "subtitle": "保存草稿 → 提交待审批 → 人工审批 → 生成待外部执行任务",
        "points": [
            "审批前核对方案版本、数量和来源",
            "操作回执须与实际方案及版本关联",
            "任务只保存在本地，没有 ERP 写入",
            "回执按任务保留，服务记录变化先核对"
        ],
        "note": "幂等键帮助防重复；服务端 expected_version 原子校验仍待完善",
        "speaker": "审批成功不表示已经调出库存。execute 只创建 draft_pending_external_execution。今日待办可以按门店、商品和编号定位；任务显示对应保存版本的业务身份。未保存人工回执可跨分类和页面继续填写，保存前读取最新状态，关联一致才清除编辑。即使人工登记完成，也不能声称外部操作由系统自动完成。读取与写入之间的原子版本检查仍由魏完善。"
    },
    {
        "eyebrow": "04 / PURCHASE SCENARIO",
        "title": "先确认减量条件\n再看成本与缺货风险",
        "subtitle": "周期、门店、品类、减量比例 → 参数预览 → 确认计算",
        "points": [
            "当前只支持减少可调整采购数量",
            "备注文字不会自动变成任意业务动作",
            "同时比较采购支出、期末库存与风险",
            "条件改变后旧预览和结果失效",
            "分周采购支出不是银行账户余额"
        ],
        "note": "少采购、库存成本减少与现金到账分别解释",
        "speaker": "演示未来十四天减少百分之二十采购。先展示参数卡，确认后才调用后端。结果除了支出变化，还必须检查新增缺货风险；不能把减量金额直接说成利润或回款。"
    },
    {
        "eyebrow": "05 / INPUT AND EVIDENCE",
        "title": "输入可以复现\n缺项需要明确保留",
        "subtitle": "后端内置种子 · 50 个门店 · 313 个库存输入 · 独立付款计划",
        "points": [
            "输入导出带来源、日期、单位与哈希",
            "下载文件不代表已经导入当前服务",
            "反馈规则草稿必须由人工核对",
            "两组各24条人工参考尚未运行模型评估"
        ],
        "note": "合成数据、人工参考与实际模型效果不能互相替代",
        "speaker": "导出包从当前后端种子读取稳定输入，在临时数据库中生成，不打开用户日常数据库。实际模型尚未接入，因此不报告模型准确率或 token 成本结论。"
    },
    {
        "eyebrow": "06 / VERIFICATION",
        "title": "让真实操作链\n经得起重复检查",
        "subtitle": "浏览器 → 前端代理 → 实际 FastAPI → 临时数据库",
        "points": [
            "覆盖工作台、审批、任务与人工回执",
            "覆盖反馈核对和事实变化后重算",
            "覆盖模拟预览、约束错误与真实缺项",
            "样例校验含零减量、全减量与数据隔离"
        ],
        "note": "本地复验：JS 138/138 · Python 163/163 · 并发读取 150/150",
        "speaker": "后端仍为魏 lanyangyang 的19794ca（并发修复3cc6c4a）。2026-10-03本地统一验收六步全部通过，Python测试包含43项实际后端浏览器用例；独立45次串行和9线程150次并发读取均为200，内容与串行基准一致。没有用固定响应替代计算，没有跳过失败用例。Windows与Ubuntu使用相同验收入口，远端结果以Actions对应提交为准；仓库中的旧联调录像未替换，本轮证据保存在验收目录。这些结果不证明模型理解、客户收益、跨进程事务或ERP执行。"
    },
    {
        "eyebrow": "07 / NEXT VALIDATION",
        "title": "下一步先验证\n能否减少实际工作量",
        "subtitle": "聚焦一条业务流程，比较完整任务耗时、可执行率和有效决策成本。",
        "points": [
            "完善跨进程事务、版本失效与权限",
            "接入一份授权真实数据，先做只读核对",
            "对比熟练人员加 GPT 的同条件基线",
            "模型按需使用，执行结果如实回收"
        ],
        "note": "多 Agent、模型接口与页面数量本身不构成商业壁垒",
        "speaker": "当前差异化还需要实验验证。应该给熟练人员和 GPT 同样的数据、约束与计算工具，比较完整任务成本。后续可能积累的是实际约束、可执行流程和结果反馈，而不是简单增加模型调用次数。"
    }
]


def build_html() -> None:
    sections = []
    for index, item in enumerate(SLIDES, 1):
        title = escape(item["title"]).replace("\n", "<br>")
        points = "".join(f"<li>{escape(point)}</li>" for point in item["points"])
        sections.append(
            f'<section class="slide" aria-label="第 {index} 页"{ "" if index == 1 else " hidden"}>'
            f'<div class="eyebrow">{escape(item["eyebrow"])}</div><h1>{title}</h1>'
            f'<p class="subtitle">{escape(item["subtitle"])}</p><ul>{points}</ul>'
            f'<footer><span>{escape(item["note"])}</span><b>{index:02d} / {len(SLIDES):02d}</b></footer>'
            f'<details><summary>讲稿</summary><p>{escape(item["speaker"])}</p></details></section>'
        )
    html = '''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>货不压钱 · 演示稿</title><style>
:root{color-scheme:dark;font-family:"Microsoft YaHei","PingFang SC",sans-serif;background:#0d1723;color:#eef3f6}
*{box-sizing:border-box}body{margin:0;padding:24px}main{max-width:1240px;margin:auto}
.slide{position:relative;min-height:697px;padding:52px 64px 70px;border:1px solid #2b3b4b;background:#132232}
.slide[hidden]{display:none}.eyebrow{color:#9ad7b7;font-size:14px;font-weight:700;letter-spacing:.15em}
h1{font-size:50px;line-height:1.25;letter-spacing:-.025em;margin:24px 0 16px;max-width:1050px}
.subtitle{color:#bdcbd5;font-size:20px;line-height:1.5;margin:0 0 25px}ul{padding:0;list-style:none;margin:0}
li{padding:9px 0 9px 23px;font-size:20px;line-height:1.4;position:relative}li:before{content:"";width:5px;height:5px;background:#9ad7b7;position:absolute;top:20px;left:0}
footer{position:absolute;bottom:30px;left:64px;right:64px;display:flex;justify-content:space-between;gap:20px;border-top:1px solid #354656;padding-top:14px;font-size:12px;color:#b8c8d2}
footer b{white-space:nowrap;color:#9ad7b7}details{margin-top:20px;color:#b8c8d2;font-size:14px}details p{max-width:900px;line-height:1.6}
nav{display:flex;justify-content:center;align-items:center;flex-wrap:wrap;gap:12px;margin:16px auto;font-size:14px}
button,a{font:inherit;border:1px solid #4e6273;color:#eef3f6;border-radius:4px;background:#132232;padding:9px 16px;text-decoration:none;cursor:pointer}button:disabled{opacity:.4;cursor:default}button:focus-visible,a:focus-visible{outline:3px solid #9ad7b7}
@media(max-width:700px){body{padding:12px}.slide{padding:30px 24px 96px;min-height:calc(100vh - 115px)}h1{font-size:32px}.subtitle,li{font-size:17px}footer{left:24px;right:24px;bottom:22px;font-size:11px}}
@media print{@page{size:13.333in 7.5in;margin:0}body{padding:0}main{max-width:none}.slide,.slide[hidden]{display:block;width:13.333in;height:7.5in;min-height:0;break-after:page;padding:50px 64px;background:#fff;color:#122333;border:0;print-color-adjust:exact}.slide:last-child{break-after:auto}.eyebrow,footer b{color:#23704f}.subtitle,footer{color:#3b5366}li:before{background:#23704f}h1{font-size:48px}li{font-size:19px}nav,details{display:none}}
</style></head><body><main>''' + "\n".join(sections) + '''</main>
<nav aria-label="演示页控制"><button id="previous" type="button">上一页</button><span id="position" aria-live="polite"></span><button id="next" type="button">下一页</button><button id="print" type="button">打印 / PDF</button><a href="DEMO_SLIDES.pptx">下载 PowerPoint</a></nav>
<script>
const slides=[...document.querySelectorAll('.slide')];let current=0;
function show(index){current=Math.max(0,Math.min(slides.length-1,index));slides.forEach((slide,i)=>{slide.hidden=i!==current});document.querySelector('#position').textContent=`${current+1} / ${slides.length}`;document.querySelector('#previous').disabled=current===0;document.querySelector('#next').disabled=current===slides.length-1;document.title=`${current+1}/${slides.length} · 货不压钱演示稿`;}
document.querySelector('#previous').addEventListener('click',()=>show(current-1));document.querySelector('#next').addEventListener('click',()=>show(current+1));document.querySelector('#print').addEventListener('click',()=>window.print());document.addEventListener('keydown',event=>{if(['ArrowRight','PageDown'].includes(event.key)){event.preventDefault();show(current+1)}else if(['ArrowLeft','PageUp'].includes(event.key)){event.preventDefault();show(current-1)}});show(0);
</script></body></html>'''
    (ROOT / "docs/DEMO_SLIDES.html").write_text(html, encoding="utf-8")


def build_powerpoint() -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "货不压钱 · v1.2 实际后端联调"
    presentation.core_properties.subject = "v1.2 / 合成零食数据 / 确定性业务计算 / zmj"
    presentation.core_properties.author = "货不压钱项目组"
    presentation.core_properties.language = "zh-CN"

    def text_box(slide, value, left, top, width, height, size, color, bold=False):
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for number, line in enumerate(value.split("\n")):
            paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(5)
        return shape

    for index, item in enumerate(SLIDES, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("132232")
        text_box(slide, item["eyebrow"], .68, .40, 12, .35, 13, "9AD7B7", True)
        text_box(slide, item["title"], .68, 1.00, 12, 1.55, 36, "EEF3F6", True)
        text_box(slide, item["subtitle"], .68, 2.70, 12, .55, 17, "BDCBD5")
        for row, point in enumerate(item["points"]):
            mark = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.70), Inches(3.50 + row * .48), Inches(.045), Inches(.045))
            mark.fill.solid()
            mark.fill.fore_color.rgb = RGBColor.from_string("9AD7B7")
            mark.line.fill.background()
            text_box(slide, point, .90, 3.37 + row * .48, 11.6, .43, 18, "EEF3F6")
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.68), Inches(6.68), Inches(11.95), Inches(.012))
        rule.fill.solid()
        rule.fill.fore_color.rgb = RGBColor.from_string("354656")
        rule.line.fill.background()
        text_box(slide, item["note"], .68, 6.86, 10.8, .35, 10, "BDCBD5")
        text_box(slide, f"{index:02d} / {len(SLIDES):02d}", 11.88, 6.86, .8, .35, 10, "9AD7B7", True)
        slide.notes_slide.notes_text_frame.text = item["speaker"]
    presentation.save(ROOT / "docs/DEMO_SLIDES.pptx")


if __name__ == "__main__":
    build_html()
    build_powerpoint()
    print(f"Built {len(SLIDES)} slides: docs/DEMO_SLIDES.html and docs/DEMO_SLIDES.pptx")
