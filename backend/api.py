"""FastAPI API 与当前静态前端之间的最小真实闭环。"""

from __future__ import annotations

import base64
import csv
import io
import json
import os
from datetime import datetime, date, timezone
from decimal import Decimal, ROUND_CEILING
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, HTTPException, Header, Query
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, ConfigDict

from .domain import (
    TEACHER_BASELINE_VERSION,
    calculate_expiry_rescue,
    calculate_procurement_brake,
    calculate_transfer,
    compare_sales,
    dec,
    dedupe_attention_cost,
    money,
    parse_feedback_dates,
    scenario_cash,
    solve_cash_goal,
    teacher_baseline,
    validate_action_bundle,
)
from .store import Store
from .errors import BusinessConflict
from .serialization import numeric_response, dumps
from .imports import inventory_rows, REQUIRED_INVENTORY_FIELDS
from .retail import demo_dataset, demo_overview, simulate_purchase, simulation_options


ROOT = Path(__file__).resolve().parents[1]
DB_PATH = os.environ.get("INVENTORY_AGENT_DB", str(ROOT / "inventory_cash_agent.db"))
store = Store(DB_PATH)
# 演示数据库使用 INSERT OR IGNORE：真实数据数据库不混入合成风险队列。
if os.environ.get("INVENTORY_AGENT_MODE", "demo") != "real":
    store.seed_demo()

app = FastAPI(title="货不压钱｜连锁零售库存资金 Agent API", version="1.3.0")


class InvestigationInput(BaseModel):
    actor_id: str = "demo-user"


class FeedbackInput(BaseModel):
    raw_text: str = Field(min_length=1)
    submitted_at: Optional[str] = None
    timezone: str = "Asia/Shanghai"
    actor_id: str = "demo-user"


class ConfirmInput(BaseModel):
    confirmed: Dict[str, Any]
    actor_id: str = "demo-user"


class RevisionInput(BaseModel):
    status: str
    actor_id: str = "demo-user"


class ReplanInput(BaseModel):
    expected_version: int = Field(ge=1, strict=True)
    simulate_failure: bool = False


class SimulationInput(BaseModel):
    target: Decimal
    horizon_days: int = Field(default=30, ge=1, le=365)
    constraints: Dict[str, Any] = Field(default_factory=dict)
    excluded_action_ids: List[str] = Field(default_factory=list)


class RetailSimulationInput(BaseModel):
    model_config = ConfigDict(extra="forbid")
    horizon_days: int = Field(default=14, ge=1, le=90, strict=True)
    reduction_pct: Decimal = Field(default=Decimal("20"), ge=0, le=100)
    store_id: Optional[str] = Field(default="all", max_length=100)
    category: Optional[str] = Field(default=None, max_length=100)
    request_text: str = Field(default="", max_length=2000)


class VersionInput(BaseModel):
    expected_version: int = Field(ge=1, strict=True)


class WorkbenchInput(BaseModel):
    expected_version: int = Field(ge=0, strict=True)
    input: Dict[str, Any] = Field(default_factory=dict)
    actor_id: str = "demo-user"


class SaveWorkbenchInput(WorkbenchInput):
    expected_proposal_version: int = Field(ge=0, strict=True)


class ExecutionStatusInput(VersionInput):
    confirmation_method: str = "manual"
    status: str
    receipt_ref: Optional[str] = None
    actual_cash: Optional[Decimal] = None
    actor_id: str = "demo-user"


class DataImportInput(BaseModel):
    filename: str
    as_of_date: Optional[date] = None
    content: str = ""
    file_base64: Optional[str] = None
    data_kind: str = "inventory"
    sheet_name: Optional[str] = None
    mode: str = "erp_file"
    actor_id: str = "demo-user"


def tenant_from_header(x_tenant_id: Optional[str]) -> str:
    return x_tenant_id or "demo"


def _service_error(exc):
    raise HTTPException(409 if isinstance(exc, BusinessConflict) else 400,
                        exc.detail() if isinstance(exc, BusinessConflict) else str(exc))


def _input_risk_id(data, config):
    value = data.get("risk_id", config["risk_id"])
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise HTTPException(422, "risk_id 必须为正整数")
    return value


WORKBENCHES = {
    "transfer": {"label": "跨门店智能调拨", "risk_type": "调拨", "risk_id": 1, "calculator": calculate_transfer},
    "expiry-rescue": {"label": "近效期现金抢救", "risk_type": "促销", "risk_id": 3, "calculator": calculate_expiry_rescue},
    "procurement-brake": {"label": "采购刹车", "risk_type": "采购刹车", "risk_id": 4, "calculator": calculate_procurement_brake},
}


from .demo_data import TRANSFER_NETWORK_STORES


def _transfer_network(input_data: Dict[str, Any]) -> List[Dict[str, Any]]:
    """50 店演示网络：评分和金额均复用调拨计算，不作为真实门店主数据。"""
    output: List[Dict[str, Any]] = []
    for district, name, store_id, on_hand, capacity, safety, daily, distance, minutes, fee, eta, demand_weight in TRANSFER_NETWORK_STORES:
        candidate_input = dict(input_data, target_store=name, target_store_id=store_id, target_on_hand=on_hand, target_capacity=capacity, target_safety=safety, target_daily_sales=daily, transport_fee=fee, eta_days=eta)
        calculation = calculate_transfer(candidate_input)
        economic = calculation.get("economic") or {}
        sale_qty = int(economic.get("target_sale_before_expiry_qty") or 0)
        score = max(0, sale_qty * 100 - fee + demand_weight)
        output.append({
            "district": district, "name": name, "store_id": store_id, "on_hand": on_hand, "capacity": capacity,
            "safety_stock": safety, "daily_sales": daily, "distance_km": distance, "travel_minutes": minutes,
            "transport_fee": fee, "eta_days": eta, "demand_score": score, "calculation": calculation,
        })
    return sorted(output, key=lambda item: item["demand_score"], reverse=True)


def _operating_store_summary(risks: List[Dict[str, Any]]) -> Dict[str, Any]:
    """总览使用的50店演示经营快照；风险字段来自当前分析快照，规模字段清晰标记为演示网络。"""
    risk_by_store = {risk["store"]: risk for risk in risks}
    stores: List[Dict[str, Any]] = []
    for index, (district, name, _store_id, on_hand, capacity, _safety, _daily, _distance, _minutes, _fee, _eta, _weight) in enumerate(TRANSFER_NETWORK_STORES):
        risk = risk_by_store.get(name)
        categories = 22 + index % 18
        inventory_value = 16800 + (capacity * 420) + (index % 7 * 1150)
        near_expiry = 0
        transfer = 0
        slow = 0
        if risk:
            inventory_value = max(inventory_value, int(float(risk["inventory_qty"] or 0) * float(risk["unit_cost"] or 0)))
            near_expiry = 1 if "near_expiry" in (risk.get("tags") or []) else 0
            transfer = 1 if risk["risk_type"] == "调拨" else 0
            slow = 1 if "slow" in (risk.get("tags") or []) else 0
        stores.append({
            "district": district, "name": name, "categories": categories,
            "inventory_value": inventory_value, "remaining_units": on_hand,
            "near_expiry_lots": near_expiry, "transfer_suggestions": transfer,
            "slow_moving_skus": slow,
        })
    total_inventory = sum(item["inventory_value"] for item in stores)
    return {
        "source_label": "50店演示经营快照",
        "season_label": "本季度（演示）",
        "season_purchase_amount": 486200,
        "store_count": len(stores),
        "category_count": 39,
        "total_inventory_value": total_inventory,
        "near_expiry_lots": sum(item["near_expiry_lots"] for item in stores),
        "transfer_suggestions": sum(item["transfer_suggestions"] for item in stores),
        "slow_moving_skus": sum(item["slow_moving_skus"] for item in stores),
        "stores": stores,
    }


def _real_operating_summary(snapshot: Dict[str, Any], tenant: str) -> Dict[str, Any]:
    """把真实库存快照映射为总览页面需要的门店经营结构。

    只映射源文件真正提供的库存和成本事实；销量、效期和调拨字段明确保持不可用。
    """
    rows = store.real_inventory_store_summary(snapshot["id"], tenant)
    prefix = "惠州市德源堂医药连锁有限公司"
    stores = []
    for row in rows:
        short_name = (row["org_name"].replace(prefix, "").replace("分店", "").strip()) or "公司主体"
        stores.append({
            "district": "惠州" if row["org_name"].startswith(prefix) else None,
            "name": short_name,
            "org_code": row["org_code"],
            "categories": row["categories"],
            "inventory_value": row["inventory_value"],
            "remaining_units": row["available_qty"],
            "near_expiry_lots": None,
            "transfer_suggestions": None,
            "slow_moving_skus": row["slow_moving_skus"],
            "stopped_purchase_value": row["stopped_purchase_value"],
            "stopped_purchase_rows": row["stopped_purchase_rows"],
            "candidate_inventory_value": row["candidate_inventory_value"],
            "suggested_reduction_amount": row["suggested_reduction_amount"],
            "p1_count": row["p1_count"],
            "p2_count": row["p2_count"],
            "p3_count": row["p3_count"],
        })
    total_value = sum(float(item["inventory_value"] or 0) for item in stores)
    metadata = snapshot.get("metadata") or {}
    return {
        "source_label": f"真实库存快照（{snapshot.get('as_of_date') or '未标注时点'}）",
        "season_label": "库存时点",
        "season_purchase_amount": None,
        "store_count": len(stores),
        "category_count": int(snapshot.get("sku_count") or 0),
        "total_inventory_value": total_value,
        "near_expiry_lots": None,
        "transfer_suggestions": None,
        "slow_moving_skus": sum(item["slow_moving_skus"] for item in stores),
        "remaining_units_label": "多单位，按商品单位分别统计",
        "region_label": f"全部{len(stores)}个门店",
        "mixed_units": True,
        "stores": stores,
        "data_quality": metadata.get("quality") or {},
        "scope_note": metadata.get("scope_note"),
        "teacher_baseline": _teacher_baseline_status(snapshot, tenant),
    }


def _teacher_baseline_status(snapshot: Optional[Dict[str, Any]] = None, tenant: str = "demo") -> Dict[str, Any]:
    """返回老师口径的输入状态与已导入快照的确定性汇总。"""
    metadata = (snapshot or {}).get("metadata") or {}
    quality = metadata.get("quality") or {}
    missing = list(metadata.get("teacher_baseline_missing_fields") or [])
    if not missing:
        summary = store.real_teacher_baseline_summary(snapshot["id"], tenant) if snapshot else {}
        return {
            "version": TEACHER_BASELINE_VERSION,
            "status": "ready",
            **summary,
            "missing_fields": [],
            "incomplete_rows": int(metadata.get("teacher_baseline_incomplete_rows") or 0),
            "source_quality": quality,
        }
    return {
        "version": TEACHER_BASELINE_VERSION,
        "status": "blocked",
        "candidate_count": 0,
        "p1_count": 0,
        "p2_count": 0,
        "p3_count": 0,
        "suggested_reduction_amount": None,
        "missing_fields": missing,
        "source_quality": quality,
    }


def _teacher_baseline_for_risk(risk: Dict[str, Any]) -> Dict[str, Any]:
    """把已有风险事实适配到老师基线；缺少扩展字段时保持 blocked。"""
    unit_cost = money(risk.get("unit_cost"))
    inventory_qty = dec(risk.get("inventory_qty"))
    inventory_amount = inventory_qty * unit_cost if inventory_qty is not None and unit_cost is not None else None
    tags = set(risk.get("tags") or [])
    return teacher_baseline(
        {
            "inventory_amount": inventory_amount,
            "available_qty": inventory_qty,
            "sales_30": risk.get("sales_30"),
            "sales_90": risk.get("sales_90"),
            "sales_cost_30": risk.get("sales_cost_30"),
            "stat_class": risk.get("stat_class"),
            "purchase_status": risk.get("purchase_status"),
            "stockout": "stockout" in tags,
            "near_stockout": "near_stockout" in tags,
        }
    )


def _decorate_teacher_risk(risk: Dict[str, Any]) -> Dict[str, Any]:
    baseline = _teacher_baseline_for_risk(risk)
    risk["teacher_baseline"] = baseline
    risk["teacher_priority"] = baseline.get("priority")
    risk["calculation_version"] = TEACHER_BASELINE_VERSION
    return risk


def _real_teacher_risk(line: Dict[str, Any], snapshot: Dict[str, Any]) -> Dict[str, Any]:
    """将已持久化的老师口径候选映射为只读诊断队列，不推断处置动作。"""
    quantity = float(line.get("inventory_qty") or 0)
    amount = float(line.get("cost_amount") or 0)
    sales_30 = float(line.get("sales_30") or 0)
    unit_cost = amount / quantity if quantity else 0
    days_to_sell = int((Decimal(str(quantity)) * Decimal("30") / Decimal(str(sales_30))).to_integral_value(rounding=ROUND_CEILING)) if sales_30 > 0 else None
    prefix = "惠州市德源堂医药连锁有限公司"
    store_name = (str(line.get("org_name") or "").replace(prefix, "").replace("分店", "").strip()) or "公司主体"
    effective = store.risk(int(line["id"]), snapshot["tenant_id"])
    if effective and effective["snapshot_id"] == snapshot["id"]:
        quantity = effective["inventory_qty"]
        sales_30 = effective["sales_30"]
        unit_cost = effective["unit_cost"]
        amount = money(quantity * unit_cost) if quantity is not None and unit_cost is not None else None
        days_to_sell = int((Decimal(str(quantity)) * 30 / Decimal(str(sales_30))).to_integral_value(rounding=ROUND_CEILING)) if quantity is not None and sales_30 else None
    else:
        effective = None
    baseline = teacher_baseline({**line, "inventory_qty": quantity, "inventory_amount": amount, "sales_30": sales_30})
    reduction = float(baseline.get("suggested_reduction_amount") or 0)
    reason = baseline.get("trigger_reason") or "老师口径候选"
    return {
        "id": int(line["id"]),
        "current_fact_version": (effective or {}).get("current_fact_version"),
        "investigation_status": (effective or {}).get("investigation_status", "not_started"),
        "fact_source": "人工确认覆盖" if effective and effective["current_fact_version"] > 1 else "文件快照",
        "source_inventory_qty": line.get("inventory_qty"),
        "snapshot_id": snapshot["id"],
        "sku": line.get("sku"),
        "product": line.get("product_name"),
        "store": store_name,
        "store_id": line.get("org_code"),
        "unit": line.get("unit"),
        "sales_30": sales_30,
        "sales_90": float(line.get("sales_90") or 0),
        "comparison": [],
        "inventory_qty": quantity,
        "unit_cost": unit_cost,
        "days_to_sell": days_to_sell,
        "risk_type": "滞销",
        "priority": line.get("teacher_priority"),
        "teacher_priority": line.get("teacher_priority"),
        "observation": f"{reason}；建议压降库存金额 ¥{reduction:,.2f}。",
        "evidence_level": "partial",
        "evidence_label": "部分支持",
        "missing_fields": [
            "shelf_availability", "store_scale", "trade_area_tag", "price_history",
            "seasonality_history", "purchase_order_status", "in_transit_qty",
            "replenishment_rule", "replenishment_log",
        ],
        "created_at": snapshot.get("created_at"),
        "teacher_baseline": baseline,
        "calculation_version": TEACHER_BASELINE_VERSION,
    }


def _workbench_config(module_type: str) -> Dict[str, Any]:
    config = WORKBENCHES.get(module_type)
    if not config:
        raise HTTPException(404, "工作台不存在")
    return config


def _risk_factors(risk: Dict[str, Any]) -> List[Dict[str, Any]]:
    """原因是假设；只有实际取得的字段才能构成支持，绝不由“未标缺项”倒推。"""
    factor_specs = [
        ("门店客群不匹配", "partial", ["本店近30天销量", "同规格门店销量对照"], ["store_scale", "trade_area_tag"], "补充门店客群和商圈标签后再判断是否为客群不匹配。"),
        ("定价不合理", "insufficient", [], ["price_history"], "补充本店价格与折扣历史后再判断。"),
        ("陈列或推荐不足", "insufficient", [], ["shelf_availability", "stockout_records"], "补充上架、陈列及缺货记录后再判断。"),
        ("季节已经过去", "insufficient", [], ["seasonality_history"], "补充去年同期和季节曲线后再判断。"),
        ("采购量过大", "insufficient", [], ["purchase_order_status", "in_transit_qty"], "补充采购单、在途明细和历史订货量后再判断。"),
        ("销售下降但仍在自动补货", "insufficient", [], ["replenishment_rule", "replenishment_log"], "补充补货规则与触发记录后再判断。"),
    ]
    output = []
    for label, level, basis, missing_fields, hint in factor_specs:
        output.append({"label": label, "kind": "原因假设", "evidence_level": level, "evidence_label": {"sufficient": "证据充分", "partial": "部分支持", "insufficient": "信息不足"}[level], "basis": basis, "missing_fields": missing_fields, "next_check": hint})
    return output


def _diagnosis_report(risk: Dict[str, Any], comparison: Dict[str, Any]) -> Dict[str, Any]:
    peer_median = comparison.get("comparison_median")
    sales = risk.get("sales_30")
    gap = None
    if peer_median not in (None, 0) and sales is not None:
        gap = round((1 - float(sales) / float(peer_median)) * 100)
    if gap is not None and gap > 0:
        conclusion = f"已确认：本店近30天销量比同规格对照门店中位数低 {gap}%，且当前库存还可覆盖约 {risk.get('days_to_sell')} 天，属于优先处理的滞销资金占用。"
    else:
        conclusion = f"已确认：当前库存覆盖约 {risk.get('days_to_sell')} 天，超过常规周转观察阈值，需优先处理库存占用。"
    return {
        "status": "可确认的异常",
        "conclusion": conclusion,
        "cannot_conclude": "现有数据只能确认销量偏低和库存覆盖偏高；尚不能确认是定价、陈列、季节、客群还是采购补货造成。",
        "suggested_next_step": "先核查门店客群与商圈标签；如确认该店需求偏低，再进入跨店调拨或采购刹车方案测算。",
    }


def _workbench_payload(module_type: str, tenant: str, risk_id: Optional[int] = None) -> Dict[str, Any]:
    config = _workbench_config(module_type)
    selected_risk_id = risk_id or config["risk_id"]
    risk = store.risk(selected_risk_id, tenant)
    if not risk:
        raise HTTPException(404, "风险不存在或无权访问")
    risk = _decorate_teacher_risk(risk)
    if risk["risk_type"] != config["risk_type"]:
        raise HTTPException(422, "该商品不属于此工作台的处置类型，请从对应业务入口打开")
    draft = store.workbench_draft(module_type, selected_risk_id, tenant)
    try:
        input_data = store.workbench_input(module_type, selected_risk_id, {}, tenant)
    except ValueError as exc:
        _service_error(exc)
    calculation = (draft or {}).get("calculation")
    # 计算口径升级后，旧草稿缺少新增结果字段时以原输入重新推演展示；不改写草稿状态或审批版本。
    if (calculation is None or (module_type == "transfer" and "economic" not in calculation)) and module_type in WORKBENCHES:
        calculation = config["calculator"](input_data)
    proposal = store.proposal((draft or {}).get("proposal_id"), tenant) if (draft or {}).get("proposal_id") else (store.proposal(risk["proposal_id"], tenant) if risk.get("proposal_id") else None)
    tasks = [task for task in store.execution_tasks(tenant) if proposal and task["proposal_id"] == proposal["id"]]
    items = [_decorate_teacher_risk(item) for item in store.risks(tenant) if item["risk_type"] == config["risk_type"]]
    payload = {
        "module_type": module_type,
        "label": config["label"],
        "mode": "sample_replay",
        "calculation_version": TEACHER_BASELINE_VERSION,
        "snapshot_id": risk["snapshot_id"],
        "items": items,
        "risk": risk,
        "input": input_data,
        "calculation": calculation,
        "draft": draft or {"status": "not_saved", "version": 0},
        "proposal": proposal,
        "tasks": tasks,
    }
    if module_type == "transfer":
        network = _transfer_network(input_data)
        payload["transfer_network"] = network
        payload["network_note"] = "原型使用50家演示门店的库存、销量和配送参数；导入ERP门店主数据后会替换为真实结果。"
    if module_type == "expiry-rescue":
        queue = []
        for item in items:
            queue_input = store.workbench_input("expiry-rescue", item["id"], {}, tenant)
            queue_calculation = calculate_expiry_rescue(queue_input)
            queue.append({"risk": item, "input": queue_input, "forecast": queue_calculation.get("forecast"), "cash": queue_calculation.get("cash")})
        payload["expiry_queue"] = sorted(queue, key=lambda row: int(row["input"].get("sellable_days") or 999))
        payload["queue_note"] = "队列按最晚处置日期排序；每一行对应一个门店＋商品＋批次。"
    return payload


def _candidate_actions(tenant: str, excluded: List[str], store_keywords: Optional[List[str]] = None) -> List[Dict[str, Any]]:
    """从当前已计算的真实工作台草稿取候选；不再内置情景模拟样例。"""
    output: List[Dict[str, Any]] = []
    rows = store.rows("SELECT * FROM workbench_drafts WHERE tenant_id=? AND status IN ('calculated','saved')", (tenant,))
    for row in rows:
        risk = store.risk(int(row["risk_id"]), tenant)
        if store_keywords and (not risk or not any(keyword in str(risk.get("store") or "") for keyword in store_keywords)):
            continue
        calculation = json.loads(row["calculation_json"]) if row.get("calculation_json") else None
        if not calculation or not calculation.get("valid"):
            continue
        action_id = row["id"]
        if action_id in excluded:
            continue
        cash = calculation.get("cash") or {}
        effect = money(cash.get("estimated_net_cash_improvement"))
        known = money(cash.get("known_cash_effect"))
        allocation = calculation.get("allocation") or {}
        output.append({
            "id": action_id,
            "module_type": row["module_type"],
            "calculation_version": TEACHER_BASELINE_VERSION,
            "risk_id": row["risk_id"],
            "label": WORKBENCHES.get(row["module_type"], {}).get("label", row["module_type"]),
            "cash_effect": effect,
            "known_cash_effect": known,
            "cash_complete": cash.get("completeness") == "complete",
            "missing_fields": cash.get("missing_fields") or [],
            "allocation": allocation,
            "input": json.loads(row["input_json"]),
            "proposal_id": row.get("proposal_id"),
        })
    return output


def _validate_candidate_bundle(candidates: List[Dict[str, Any]]) -> Dict[str, Any]:
    """把共享批次的占用放回组合求解校验，采购动作没有批次时不伪造库存占用。"""
    actions: List[Dict[str, Any]] = []
    available: Dict[str, int] = {}
    routes: List[tuple[str, str]] = []
    for item in candidates:
        input_data, allocation = item.get("input") or {}, item.get("allocation") or {}
        lot_id = allocation.get("lot_id")
        if not lot_id:
            continue
        quantity = int(allocation.get("quantity") or 0)
        module = item.get("module_type")
        action: Dict[str, Any] = {"lot_id": lot_id, "quantity": quantity, "type": "expiry_rescue"}
        if module == "transfer":
            action.update({"type": "transfer", "source_store_id": input_data.get("source_store_id"), "target_store_id": input_data.get("target_store_id")})
            routes.append((input_data.get("source_store_id"), input_data.get("target_store_id")))
            amount = max(0, int(input_data.get("source_on_hand") or 0) - int(input_data.get("source_safety") or 0))
        else:
            amount = int(input_data.get("inventory_qty") or 0)
        available[lot_id] = min(available.get(lot_id, amount), amount) if lot_id in available else amount
        actions.append(action)
    return validate_action_bundle(actions, available_by_lot=available, allowed_routes=routes)


def _approval_summary(proposal: Dict[str, Any], risk: Optional[Dict[str, Any]], tenant: str) -> Dict[str, str]:
    """把内部 proposal 版本翻译成负责人可判断的业务审批事项。"""
    payload = ((proposal.get("version") or {}).get("payload") or {})
    proposal_type = payload.get("proposal_type") or ""
    input_data = payload.get("input") or {}
    calculation = payload.get("calculation") or {}
    action = (payload.get("actions") or [{}])[0]
    product = (risk or {}).get("product") or input_data.get("product") or "该商品"
    source = (risk or {}).get("store") or input_data.get("source_store") or "调出门店"
    stores = {row["store_id"]: row["store"] for row in store.risks(tenant)}
    target = input_data.get("target_store") or stores.get(action.get("target_store_id")) or "接收门店"
    quantity = action.get("quantity") or (calculation.get("allocation") or {}).get("quantity") or input_data.get("quantity") or "待确认"
    if proposal_type == "transfer":
        return {"title": "审批跨店调拨：%s %s 件" % (product, quantity), "detail": "确认从%s调往%s的数量、安全库存与运输安排。" % (source, target), "boundary": "审批后仅生成待出库任务，不会自动向 ERP 发单或视为已完成调拨。"}
    if proposal_type == "expiry-rescue":
        forecast = calculation.get("forecast") or {}
        return {"title": "审批近效期处置：%s" % product, "detail": "确认调拨%s件、促销%s件、退供%s件的组合处置。" % (forecast.get("transfer_qty", "待确认"), forecast.get("promo_qty", "待确认"), forecast.get("return_qty", "待确认")), "boundary": "商品促销与沟通文案仍需人工审核，不会自动发送或改价。"}
    if proposal_type == "procurement-brake":
        order = calculation.get("order") or {}
        return {"title": "审批采购调整：%s" % product, "detail": "确认采购单%s的减量、取消或延期安排及其付款影响。" % (order.get("po_number") or input_data.get("po_number") or "待确认"), "boundary": "审批后只生成待执行方案，不会直接修改供应商订单。"}
    return {"title": "审批处置方案：%s" % product, "detail": "确认当前事实与测算版本后再进入执行。", "boundary": "审批绑定当前版本；事实或输入变更后需重新确认。"}


@app.get("/api/v1/health")
@numeric_response
def health() -> Dict[str, Any]:
    return {"status": "ok", "version": app.version}


@app.get("/api/v1/dashboard")
@numeric_response
def dashboard(x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    real_snapshot = store.real_inventory_snapshot(tenant)
    if real_snapshot is not None:
        operating_summary = _real_operating_summary(real_snapshot, tenant)
        total_inventory = float(real_snapshot.get("cost_total") or 0)
        stopped_purchase = sum(float(item.get("stopped_purchase_value") or 0) for item in operating_summary["stores"])
        teacher_status = operating_summary["teacher_baseline"]
        candidate_inventory = teacher_status.get("candidate_inventory_amount") if teacher_status.get("status") == "ready" else None
        return {
            "tenant_id": tenant,
            "calculation_version": TEACHER_BASELINE_VERSION,
            "current_inventory_cost": round(total_inventory, 2),
            "attention_inventory_cost": candidate_inventory,
            "attention_inventory_cost_wan": None if candidate_inventory is None else round(float(candidate_inventory) / 10000, 4),
            "expected_net_cash_improvement": None,
            "calculable_cash_improvement": 0,
            "suggested_inventory_reduction_amount": teacher_status.get("suggested_reduction_amount"),
            "cash_completeness": "unavailable",
            "expected_avoided_loss": None,
            "actual_execution_result": None,
            "pending_investigations": 0,
            "pending_approvals": 0,
            "priority_actions": [],
            "snapshot_id": real_snapshot["id"],
            "source_label": operating_summary["source_label"],
            "source_scope_note": operating_summary["scope_note"],
            "procurement_stop_value": round(stopped_purchase, 2),
            "teacher_baseline": teacher_status,
            "operating_summary": operating_summary,
        }
    risks = store.risks(tenant)
    lots = [{"lot_id": "LOT-RISK-%s" % r["id"], "quantity": r["inventory_qty"], "unit_cost": r["unit_cost"], "labels": r["tags"]} for r in risks]
    attention_result = dedupe_attention_cost(lots)
    total_inventory = sum((float(r["inventory_qty"] or 0) * float(r["unit_cost"] or 0) for r in risks), 0.0)
    attention = float(attention_result["amount"] or 0)
    partial = sum(1 for r in risks if r["evidence_level"] != "sufficient")
    pending = len(store.rows("SELECT id FROM proposals WHERE tenant_id=? AND status IN ('pending_approval','needs_replan')", (tenant,)))
    tasks = store.execution_tasks(tenant)
    actual = sum(float(task["metadata"].get("actual_cash") or 0) for task in tasks if task["status"] in ("received", "completed")) or None
    candidates = _candidate_actions(tenant, [])
    known_total = sum(float(item["cash_effect"] or 0) for item in candidates if item["cash_effect"] is not None)
    unavailable = any(item["cash_effect"] is None for item in candidates)
    avoided = sum(float((json.loads(row["calculation_json"]).get("cash") or {}).get("avoided_loss") or 0) for row in store.rows("SELECT calculation_json FROM workbench_drafts WHERE tenant_id=? AND calculation_json IS NOT NULL", (tenant,))) or None
    return {"tenant_id": tenant, "calculation_version": TEACHER_BASELINE_VERSION, "current_inventory_cost": round(total_inventory, 2), "attention_inventory_cost": round(attention, 2), "attention_inventory_cost_wan": round(attention / 10000, 4), "expected_net_cash_improvement": None if unavailable else round(known_total, 2), "calculable_cash_improvement": round(known_total, 2), "cash_completeness": "unavailable" if unavailable else "complete", "expected_avoided_loss": avoided, "actual_execution_result": actual, "pending_investigations": partial, "pending_approvals": pending, "priority_actions": candidates[:5], "snapshot_id": "snapshot-demo-v1", "source_label": "合成样例回放", "teacher_baseline": {"version": TEACHER_BASELINE_VERSION, "status": "sample_replay", "candidate_count": 0, "p1_count": 0, "p2_count": 0, "p3_count": 0, "suggested_reduction_amount": None, "missing_fields": ["sales_90", "sales_cost_30", "stat_class", "stockout", "near_stockout"]}, "operating_summary": _operating_store_summary(risks)}


@app.get("/api/v1/risks")
@numeric_response
def list_risks(
    priority: Optional[str] = Query(default=None, pattern="^P[123]$"),
    x_tenant_id: Optional[str] = Header(default=None),
) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    snapshot = store.real_inventory_snapshot(tenant)
    if snapshot is not None:
        summary = _teacher_baseline_status(snapshot, tenant)
        lines = store.real_teacher_candidates(snapshot["id"], tenant, limit=300, priority=priority)
        filtered_total = summary.get(f"{priority.lower()}_count") if priority else summary.get("candidate_count", 0)
        return {
            "items": [_real_teacher_risk(line, snapshot) for line in lines],
            "total": summary.get("candidate_count", 0),
            "filtered_total": filtered_total,
            "display_limit": 300,
            "calculation_version": TEACHER_BASELINE_VERSION,
        }
    items = [_decorate_teacher_risk(item) for item in store.risks(tenant)]
    return {"items": items, "total": len(items)}


@app.get("/api/v1/risks/{risk_id}")
@numeric_response
def get_risk(risk_id: int, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    real_snapshot = store.real_inventory_snapshot(tenant)
    if real_snapshot is not None:
        line = store.real_teacher_line(real_snapshot["id"], risk_id, tenant)
        if not line:
            raise HTTPException(404, "老师口径候选不存在或不在当前快照")
        risk = _real_teacher_risk(line, real_snapshot)
        comparison_values = store.real_sku_sales_comparison(real_snapshot["id"], str(line["sku"]), str(line["org_code"]), tenant)
        comparison = compare_sales(risk["sales_30"], [int(value) for value in comparison_values])
        risk["comparison"] = comparison_values
        facts = [
            {"label": "当前库存", "value": f"{risk['inventory_qty']} {risk.get('unit') or '件'}", "source": risk["fact_source"], "source_detail": f"{real_snapshot['id']} · SKU＋门店库存 · 事实版本 {risk.get('current_fact_version')}", "as_of": real_snapshot.get("as_of_date")},
            {"label": "近30天销量", "value": f"{risk['sales_30']} {risk.get('unit') or '件'}", "source": "源表汇总", "source_detail": "源表三十天销量", "as_of": real_snapshot.get("as_of_date")},
            {"label": "近90天销量", "value": f"{risk['sales_90']} {risk.get('unit') or '件'}", "source": "源表汇总", "source_detail": "源表九十天销量", "as_of": real_snapshot.get("as_of_date")},
            {"label": "老师口径存销比", "value": str(risk["teacher_baseline"].get("teacher_ratio")), "source": "系统重算", "source_detail": "库存数量 ÷ MAX(30天销量, 向上取整后的月均销量)", "as_of": real_snapshot.get("as_of_date")},
            {"label": "降库存存销比", "value": str(risk["teacher_baseline"].get("reduction_ratio")), "source": "系统计算", "source_detail": "库存金额 ÷ 三十天成本", "as_of": real_snapshot.get("as_of_date")},
            {"label": "建议压降金额", "value": f"¥{float(risk['teacher_baseline'].get('suggested_reduction_amount') or 0):,.2f}", "source": "系统计算", "source_detail": "MAX(库存金额 − 三十天成本 × 3, 0)", "as_of": real_snapshot.get("as_of_date")},
        ]
        return {
            "risk": risk,
            "comparison": comparison,
            "facts": facts,
            "diagnosis": _diagnosis_report(risk, comparison),
            "factors": _risk_factors(risk),
            "teacher_baseline": risk["teacher_baseline"],
            "calculation_version": TEACHER_BASELINE_VERSION,
            "evidence": {"level": "partial", "label": "部分支持", "missing_fields": risk["missing_fields"], "supporting_refs": ["inventory_qty", "sales_30", "sales_90", "sales_cost_30", "stat_class"], "conflicts": []},
        }
    risk = store.risk(risk_id, tenant)
    if not risk:
        raise HTTPException(404, "风险不存在或无权访问")
    comparison = compare_sales(risk["sales_30"], risk["comparison"])
    snapshot = risk.get("snapshot_id") or "snapshot-demo-v1"
    facts = [
        {"label": "当前库存", "value": f"{risk['inventory_qty']} 件", "source": "库存快照", "source_detail": f"{snapshot} · SKU＋门店库存", "as_of": risk.get("created_at")},
        {"label": "近30天销量", "value": f"{risk['sales_30']} 件", "source": "销售明细汇总", "source_detail": f"{snapshot} · 截止当前快照往前30天", "as_of": risk.get("created_at")},
        {"label": "对照门店销量中位数", "value": f"{comparison.get('comparison_median')} 件" if comparison.get("comparison_median") is not None else "未知", "source": "同规格门店销售对照", "source_detail": f"{snapshot} · {comparison.get('comparison_count', 0)} 家对照门店", "as_of": risk.get("created_at")},
        {"label": "库存覆盖天数", "value": f"{risk['days_to_sell']} 天" if risk.get("days_to_sell") is not None else "未知", "source": "系统计算", "source_detail": "当前库存 ÷ 近30天日均销量", "as_of": risk.get("created_at")},
    ]
    return {"risk": risk, "comparison": comparison, "facts": facts, "diagnosis": _diagnosis_report(risk, comparison), "factors": _risk_factors(risk), "teacher_baseline": _teacher_baseline_for_risk(risk), "calculation_version": TEACHER_BASELINE_VERSION, "evidence": {"level": "partial", "label": "部分支持", "missing_fields": risk["missing_fields"], "supporting_refs": ["inventory_qty", "sales_30", "comparison_sales"], "conflicts": []}}


@app.post("/api/v1/risks/{risk_id}/investigations")
@numeric_response
def create_investigation(risk_id: int, payload: InvestigationInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    if not store.risk(risk_id, tenant):
        raise HTTPException(404, "风险不存在或无权访问")
    return store.create_investigation(risk_id, payload.actor_id, tenant)


@app.post("/api/v1/investigations/{investigation_id}/feedback")
@numeric_response
def create_feedback(investigation_id: str, payload: FeedbackInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    try:
        submitted = datetime.fromisoformat((payload.submitted_at or datetime.now(timezone.utc).isoformat()).replace("Z", "+00:00"))
        if submitted.tzinfo is None:
            raise ValueError("提交时间必须带时区")
        dates = parse_feedback_dates(payload.raw_text, submitted, payload.timezone)
    except (ValueError, KeyError) as exc:
        raise HTTPException(422, "时间或时区不合法：" + str(exc))
    draft = {"factors": [], "current_status": "unknown", "extraction_status": "manual_required", "missing_fields": ["verified_factors"], "date_interpretation": dates, "raw_text": payload.raw_text, "source": "负责人输入", "evidence_refs": [], "remediation_status": "unknown"}
    try:
        return store.add_feedback(investigation_id, payload.raw_text, submitted.isoformat(), draft, payload.actor_id, tenant)
    except ValueError as exc:
        _service_error(exc)


@app.post("/api/v1/feedback/{feedback_id}/confirm")
@numeric_response
def confirm_feedback(feedback_id: str, payload: ConfirmInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    try:
        return store.confirm_feedback(feedback_id, payload.confirmed, payload.actor_id, tenant_from_header(x_tenant_id))
    except ValueError as exc:
        raise HTTPException(409 if isinstance(exc, BusinessConflict) else 400, exc.detail() if isinstance(exc, BusinessConflict) else str(exc))


@app.post("/api/v1/feedback/{feedback_id}/revisions")
@numeric_response
def revise_feedback(feedback_id: str, payload: RevisionInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    try:
        return store.revise_feedback(feedback_id, payload.status, payload.actor_id, tenant_from_header(x_tenant_id))
    except ValueError as exc:
        raise HTTPException(409 if isinstance(exc, BusinessConflict) else 400, exc.detail() if isinstance(exc, BusinessConflict) else str(exc))


@app.post("/api/v1/risks/{risk_id}/replan")
@numeric_response
def replan(risk_id: int, payload: ReplanInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    if not store.risk(risk_id, tenant):
        raise HTTPException(404, "风险不存在或无权访问")
    proposal = store.one("SELECT status FROM proposals WHERE risk_id=? AND tenant_id=? ORDER BY created_at DESC LIMIT 1", (risk_id, tenant))
    if not proposal or proposal["status"] not in ("needs_replan", "replan_pending"):
        raise HTTPException(409, "没有因新事实而待重算的方案")
    try:
        if payload.simulate_failure:
            return store.mark_replan_pending(risk_id, tenant, expected_version=payload.expected_version)
        return store.replan(risk_id, tenant, expected_version=payload.expected_version)
    except ValueError as exc:
        if not isinstance(exc, BusinessConflict) or exc.code not in {"version_conflict", "invalid_state"}:
            try:
                store.mark_replan_pending(risk_id, tenant, str(exc), expected_version=payload.expected_version)
            except BusinessConflict:
                pass  # A concurrent successful replan must not be invalidated.
        raise HTTPException(409 if isinstance(exc, BusinessConflict) else 400, exc.detail() if isinstance(exc, BusinessConflict) else str(exc))


@app.get("/api/v1/workbenches/{module_type}")
@numeric_response
def get_workbench(module_type: str, risk_id: Optional[int] = Query(default=None), x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    return _workbench_payload(module_type, tenant_from_header(x_tenant_id), risk_id)


@app.post("/api/v1/workbenches/{module_type}/draft")
@numeric_response
def update_workbench_draft(module_type: str, payload: WorkbenchInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    config = _workbench_config(module_type)
    risk_id = _input_risk_id(payload.input, config)
    try:
        result = store.prepare_workbench(module_type, risk_id, payload.input, payload.expected_version, payload.actor_id, tenant)
        return {"draft": result["draft"]}
    except ValueError as exc:
        _service_error(exc)


@app.post("/api/v1/workbenches/{module_type}/calculate")
@numeric_response
def calculate_workbench(module_type: str, payload: WorkbenchInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    config = _workbench_config(module_type)
    risk_id = _input_risk_id(payload.input, config)
    try:
        return store.prepare_workbench(module_type, risk_id, payload.input, payload.expected_version, payload.actor_id, tenant, calculate=True)
    except ValueError as exc:
        _service_error(exc)


@app.post("/api/v1/workbenches/{module_type}/save")
@numeric_response
def save_workbench(module_type: str, payload: SaveWorkbenchInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    config = _workbench_config(module_type)
    risk_id = _input_risk_id(payload.input, config)
    try:
        return store.save_workbench(module_type, risk_id, payload.input, payload.expected_version, payload.expected_proposal_version, payload.actor_id, tenant)
    except ValueError as exc:
        _service_error(exc)


@app.post("/api/v1/proposals/{proposal_id}/submit")
@numeric_response
def submit_proposal(proposal_id: str, payload: VersionInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    try:
        return store.submit_proposal(proposal_id, tenant_id=tenant_from_header(x_tenant_id), expected_version=payload.expected_version)
    except ValueError as exc:
        _service_error(exc)


@app.get("/api/v1/proposals")
@numeric_response
def list_proposals(status: Optional[str] = Query(default=None), x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    """统一方案列表，审批页与工作台共用同一方案版本。"""
    tenant = tenant_from_header(x_tenant_id)
    sql = "SELECT id FROM proposals WHERE tenant_id=?"
    params: List[Any] = [tenant]
    if status:
        sql += " AND status=?"
        params.append(status)
    sql += " ORDER BY created_at DESC"
    items = []
    for row in store.rows(sql, tuple(params)):
        proposal = store.proposal(row["id"], tenant)
        if proposal:
            proposal["approval_summary"] = _approval_summary(proposal, store.risk(proposal["risk_id"], tenant), tenant)
            items.append(proposal)
    return {"items": items}


@app.get("/api/v1/execution-tasks")
@numeric_response
def list_execution_tasks(x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    return {"items": store.execution_tasks(tenant_from_header(x_tenant_id))}


@app.post("/api/v1/execution-tasks/{task_id}/status")
@numeric_response
def update_execution_task(task_id: str, payload: ExecutionStatusInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    try:
        return store.update_execution_status(task_id, payload.status, payload.receipt_ref, payload.actual_cash, payload.actor_id, tenant_from_header(x_tenant_id), payload.expected_version, payload.confirmation_method)
    except ValueError as exc:
        raise HTTPException(409 if isinstance(exc, BusinessConflict) else 400, exc.detail() if isinstance(exc, BusinessConflict) else str(exc))


@app.get("/api/v1/proposals/{proposal_id}/versions")
@numeric_response
def proposal_versions(proposal_id: str, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    if not store.proposal(proposal_id, tenant_from_header(x_tenant_id)):
        raise HTTPException(404, "方案不存在或无权访问")
    return {"items": store.proposal_versions(proposal_id, tenant_from_header(x_tenant_id))}


@app.get("/api/v1/proposals/{proposal_id}/export")
@numeric_response
def export_proposal(proposal_id: str, version: int = Query(ge=1), x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    proposal = store.proposal(proposal_id, tenant)
    if not proposal:
        raise HTTPException(404, "方案不存在或无权访问")
    saved = next((row for row in store.proposal_versions(proposal_id, tenant) if row["version"] == version), None)
    if not saved:
        raise HTTPException(404, "方案版本不存在")
    archived = json.loads(saved["payload_json"])
    document = {"document_type": "建议单／待原系统执行", "external_write": False,
            "proposal_id": proposal_id, "proposal_version": version,
            "snapshot_id": (archived.get("basis") or {}).get("snapshot_id", proposal["snapshot_id"]), "status": saved["status"],
            "created_at": saved["created_at"], "payload": archived}
    return Response(dumps(document), media_type="application/json", headers={"Content-Disposition": f'attachment; filename="{proposal_id}-v{version}.json"'})


@app.post("/api/v1/proposals/{proposal_id}/approve")
@numeric_response
def approve_proposal(proposal_id: str, payload: VersionInput, x_tenant_id: Optional[str] = Header(default=None), idempotency_key: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    try:
        return store.approve(proposal_id, tenant_id=tenant_from_header(x_tenant_id), idem=idempotency_key, expected_version=payload.expected_version)
    except ValueError as exc:
        raise HTTPException(409 if isinstance(exc, BusinessConflict) else 400, exc.detail() if isinstance(exc, BusinessConflict) else str(exc))


@app.post("/api/v1/proposals/{proposal_id}/execute")
@numeric_response
def execute_proposal(proposal_id: str, payload: VersionInput, x_tenant_id: Optional[str] = Header(default=None), idempotency_key: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    try:
        task = store.execute(proposal_id, tenant_id=tenant_from_header(x_tenant_id), idem=idempotency_key, expected_version=payload.expected_version)
        task["metadata"] = json.loads(task.pop("metadata_json"))
        return task
    except ValueError as exc:
        raise HTTPException(409 if isinstance(exc, BusinessConflict) else 400, exc.detail() if isinstance(exc, BusinessConflict) else str(exc))


@app.post("/api/v1/scenarios/simulate")
@numeric_response
def simulate(payload: SimulationInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    candidates = _candidate_actions(tenant, payload.excluded_action_ids, payload.constraints.get("region_store_keywords"))
    constraint_text = str(payload.constraints.get("text", ""))
    if "不要降价" in constraint_text or "不促销" in constraint_text:
        candidates = [item for item in candidates if item["module_type"] != "expiry-rescue"]
    if "不调拨" in constraint_text:
        candidates = [item for item in candidates if item["module_type"] != "transfer"]
    result = solve_cash_goal(payload.target, payload.horizon_days, candidates)
    bundle = _validate_candidate_bundle(result["selected"])
    if not bundle["valid"]:
        result["selected"] = []
        result["achieved"] = Decimal("0")
        result["gap"] = money(payload.target)
        result["status"] = "gap"
        result["solver_status"] = "blocked_by_shared_inventory"
    result["constraints"] = payload.constraints
    result["bundle_validation"] = bundle
    calculable_candidates = [item for item in candidates if (money(item.get("cash_effect")) or Decimal("0")) > 0]
    result["cash_basis"] = {"baseline_id": "baseline-demo-v1", "snapshot_id": "snapshot-demo-v1", "completeness": "partial" if calculable_candidates else "unavailable", "missing_fields": ([] if calculable_candidates else (["calculated_workbench_candidates"] if not candidates else ["cash_effect_for_current_candidates"])) + (bundle["errors"] if not bundle["valid"] else []), "assumptions": ["仅组合当前已计算且未排除的工作台草稿；同一草稿只计一次", "共享批次按占用量与允许调拨路线统一校验"]}
    return result


@app.get("/api/v1/work-items")
@numeric_response
def work_items(x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    """原型中的今日工作以业务动作为中心；证据缺项只留在详情核查。"""
    tenant = tenant_from_header(x_tenant_id)
    items = []
    snapshot = store.real_inventory_snapshot(tenant)
    if snapshot:
        for line in store.real_teacher_candidates(snapshot["id"], tenant):
            risk = _real_teacher_risk(line, snapshot)
            items.append({"id": "investigate-%s" % risk["id"], "type": "库存核查", "risk_id": risk["id"],
                          "title": "%s · %s 核查库存压货原因" % (risk["store"], risk["product"]),
                          "owner": None, "due_date": None, "status": risk["investigation_status"],
                          "priority": risk["teacher_priority"], "reason": risk["observation"],
                          "route": "slow-diagnosis", "action_label": "查看证据并核查", "snapshot_id": snapshot["id"],
                          "current_fact_version": risk["current_fact_version"], "rank": len(items)+1})
        return {"items": items, "mode": "real_inventory_snapshot", "snapshot_id": snapshot["id"],
                "missing_fields": ["assigned_owner", "due_date", "verified_cause", "executable_business_rules"]}
    playbook = {
        "调拨": {"type": "跨店调拨", "route": "transfer", "time": "今天 09:45", "action": "查看调拨方案"},
        "促销": {"type": "近效期处置", "route": "expiry-rescue", "time": "今天 10:30", "action": "选择处置方案"},
        "采购刹车": {"type": "采购调整", "route": "procurement-brake", "time": "今天 11:30", "action": "查看采购建议"},
        "退供": {"type": "供应商退供", "route": "slow-diagnosis", "time": "今天 14:00", "action": "确认退供条件"},
    }
    for risk in store.risks(tenant):
        guide = playbook.get(risk["risk_type"])
        if not guide:
            continue
        title, reason = _work_item_copy(risk)
        inventory_cost = float(risk["inventory_qty"] or 0) * float(risk["unit_cost"] or 0)
        items.append({
            "id": "action-%s" % risk["id"], "type": guide["type"], "risk_id": risk["id"],
            "workbench_risk_id": risk["id"],
            "title": title, "owner": "张经理", "due_date": guide["time"],
            "status": "待处理", "priority": risk["priority"], "reason": reason, "route": guide["route"], "action_label": guide["action"],
            "impact": "库存占用 ¥%s（%s 件 × ¥%s/件）" % (format(inventory_cost, ",.0f"), risk["inventory_qty"], format(float(risk["unit_cost"] or 0), ",.0f")),
        })
    for proposal_row in store.rows("SELECT id FROM proposals WHERE tenant_id=? AND status='pending_approval'", (tenant,)):
        proposal = store.proposal(proposal_row["id"], tenant)
        if not proposal:
            continue
        summary = _approval_summary(proposal, store.risk(proposal["risk_id"], tenant), tenant)
        risk = store.risk(proposal["risk_id"], tenant)
        proposal_title = "%s · %s 调拨方案待审批" % (risk["store"], risk["product"]) if risk else summary["title"]
        items.append({"id": "approve-%s" % proposal["id"], "type": "调拨审批", "risk_id": proposal["risk_id"], "proposal_id": proposal["id"], "title": proposal_title, "owner": "张经理", "due_date": "今天 15:30", "status": "待审批", "priority": "高", "reason": summary["detail"], "boundary": summary["boundary"], "route": "approvals", "action_label": "查看并审批"})
    for task in store.execution_tasks(tenant):
        if task["status"] not in ("received", "completed"):
            items.append({"id": "execute-%s" % task["id"], "type": "执行跟进", "risk_id": None, "task_id": task["id"], "title": "跟进执行任务 %s" % task["id"], "owner": "张经理", "due_date": "今天 16:30", "status": task["status"], "priority": "高", "reason": "等待门店出库、收货或异常回执", "route": "execution", "action_label": "更新执行状态"})
    priority_order = {"紧急": 0, "高": 1, "中": 2, "低": 3}
    items.sort(key=lambda item: (priority_order.get(item["priority"], 9), item["due_date"]))
    for rank, item in enumerate(items, start=1):
        item["rank"] = rank
    return {"items": items, "mode": "prototype_playbook", "missing_fields": []}


def _work_item_copy(risk: Dict[str, Any]) -> tuple[str, str]:
    """将风险记录翻译成负责人能直接执行的今日事项。"""
    store_name = risk["store"]
    product = risk["product"]
    quantity = risk["inventory_qty"] or 0
    sales_days = risk["days_to_sell"] or "较长"
    if risk["risk_type"] == "调拨":
        return (
            "%s · %s 滞销待调拨" % (store_name, product),
            "现有库存 %s 件，按当前销量预计需 %s 天售完；已发现其他门店需求更高，请确认调出与接收门店。" % (quantity, sales_days),
        )
    if risk["risk_type"] == "促销":
        return (
            "%s · %s 近效期待处置" % (store_name, product),
            "现有库存 %s 件，按当前销量预计需 %s 天售完，已进入近效期处置窗口；今天确定调拨、促销或退供的数量。" % (quantity, sales_days),
        )
    if risk["risk_type"] == "采购刹车":
        return (
            "%s · %s 待暂停或减量采购" % (store_name, product),
            "现有库存 %s 件，近 30 天仅销售 %s 件，且仍有在途或未执行采购；先确认减量或延期，避免库存继续累积。" % (quantity, risk["sales_30"] or 0),
        )
    return (
        "%s · %s 待向供应商申请退供" % (store_name, product),
        "现有库存 %s 件，按当前销量预计需 %s 天售完；请确认供应商退换条件、可退数量和截止日期。" % (quantity, sales_days),
    )


@app.get("/api/v1/data-center")
@numeric_response
def data_center(x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    real_snapshot = store.real_inventory_snapshot(tenant)
    if real_snapshot is not None:
        metadata = real_snapshot.get("metadata") or {}
        teacher_status = _teacher_baseline_status(real_snapshot, tenant)
        has_teacher_inputs = teacher_status["status"] == "ready"
        return {
            "mode": "real_inventory_snapshot",
            "snapshot_id": real_snapshot["id"],
            "calculation_version": TEACHER_BASELINE_VERSION,
            "teacher_baseline": teacher_status,
            "sources": [{
                "name": real_snapshot["source"],
                "fields": (
                    ["库存数量", "库存金额", "30天销量", "90天销量", "30天成本", "采购状态", "统计标记类", "门店与商品主键"]
                    if has_teacher_inputs
                    else ["库存数量", "可用数量", "进价金额", "采购状态", "门店与商品主键"]
                ),
                "last_snapshot": real_snapshot["as_of_date"],
                "status": "真实库存快照",
                "missing_fields": (
                    ["批次有效期", "采购订单", "在途数量", "陈列记录", "补货规则与日志", "价格与折扣历史"]
                    if has_teacher_inputs
                    else ["销售历史", "30天销售成本", "90天销量", "老师分类", "批次有效期", "采购订单", "门店地理位置"]
                ),
                "summary": {
                    "rows": real_snapshot["record_count"],
                    "stores": real_snapshot["org_count"],
                    "skus": real_snapshot["sku_count"],
                    "cost_total": real_snapshot["cost_total"],
                    "untaxed_cost_total": real_snapshot["untaxed_cost_total"],
                    "quality": metadata.get("quality") or {},
                    "teacher_baseline": teacher_status,
                },
            }],
            "import_schemas": _import_schemas(),
            "imports": store.data_imports(tenant),
        }
    if tenant != "demo" or os.environ.get("INVENTORY_AGENT_MODE", "demo") == "real":
        return {"mode": "unavailable", "snapshot_id": None, "sources": [], "import_schemas": _import_schemas(),
                "imports": store.data_imports(tenant), "missing_fields": ["inventory_snapshot"]}
    return {"mode": "sample_replay", "snapshot_id": "snapshot-demo-v1", "calculation_version": TEACHER_BASELINE_VERSION, "teacher_baseline": {"version": TEACHER_BASELINE_VERSION, "status": "sample_replay", "candidate_count": 0, "p1_count": 0, "p2_count": 0, "p3_count": 0, "suggested_reduction_amount": None, "missing_fields": ["sales_90", "sales_cost_30", "stat_class", "stockout", "near_stockout"]}, "sources": [{"name": "当前分析快照（演示）", "fields": ["销售", "库存", "采购", "效期", "执行回执"], "last_snapshot": "snapshot-demo-v1", "status": "演示数据", "missing_fields": []}], "import_schemas": _import_schemas(), "imports": store.data_imports(tenant)}


def _import_schemas() -> Dict[str, Dict[str, Any]]:
    return {
        "inventory": {"label": "库存表", "required": sorted(REQUIRED_INVENTORY_FIELDS), "description": "每行一个门店＋SKU，必须提供单位及成本；销量、销售成本和分类用于老师口径，缺少时保持未知。"},
        "sales": {"label": "销售表", "required": ["sku", "store", "sales_qty"], "description": "每行一个商品在一个门店一个统计周期的销量；老师口径 V1 还需要统计周期、90天累计销量和销售成本。"},
        "purchase": {"label": "采购表", "required": ["sku", "store", "po_number", "open_purchase_qty"], "description": "包含未执行采购量或在途数量的采购明细。"},
        "expiry": {"label": "效期批次表", "required": ["sku", "store", "batch", "expiry_date"], "description": "每行一个商品批次及其到期日。"},
    }


def _read_import_rows(payload: DataImportInput) -> tuple[List[Dict[str, Any]], List[str], Optional[str]]:
    """将 CSV/TSV/XLSX 统一为行字典；不把上传文件写入服务端目录。"""
    extension = Path(payload.filename).suffix.lower()
    try:
        if extension not in {".csv", ".tsv", ".xlsx"}:
            raise ValueError("仅支持 CSV、TSV、XLSX")
        if len(payload.content.encode()) > 5_000_000 or len(payload.file_base64 or "") > 7_000_000:
            raise ValueError("文件超过 5MB 限制")
        if payload.file_base64:
            raw = base64.b64decode(payload.file_base64, validate=True)
            if len(raw) > 5_000_000:
                raise ValueError("文件超过 5MB 限制")
            if extension == ".xlsx":
                from openpyxl import load_workbook

                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                if payload.sheet_name and payload.sheet_name not in workbook.sheetnames:
                    raise ValueError("指定工作表不存在")
                sheet = workbook[payload.sheet_name] if payload.sheet_name else workbook.active
                values = list(sheet.iter_rows(values_only=True))
                if not values:
                    return [], [], None
                headers = [str(value).strip() if value is not None else "" for value in values[0]]
                named = [header for header in headers if header]
                if len(named) != len(set(named)):
                    raise ValueError("存在重复列名")
                rows = [{headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]} for row in values[1:] if any(value not in (None, "") for value in row)]
                workbook.close()
                return rows, headers, None
            text = raw.decode("utf-8-sig")
        else:
            text = payload.content
        delimiter = "\t" if extension == ".tsv" else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        headers = reader.fieldnames or []
        if len(headers) != len(set(headers)) or any(not header for header in headers):
            raise ValueError("列名不能为空或重复")
        rows = list(reader)
        if any(None in row for row in rows):
            raise ValueError("数据列数超过表头列数")
        return rows, headers, None
    except Exception as exc:
        return [], [], str(exc)


@app.post("/api/v1/data-center/imports")
@numeric_response
def import_data(payload: DataImportInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    schemas = _import_schemas()
    schema = schemas.get(payload.data_kind)
    if not schema:
        raise HTTPException(422, "不支持的数据类型")
    rows, header_list, parse_error = _read_import_rows(payload)
    headers = set(header_list)
    required = REQUIRED_INVENTORY_FIELDS if payload.data_kind == "inventory" else set(schema["required"])
    errors = ([{"row": 0, "message": f"文件解析失败：{parse_error}"}] if parse_error else [])
    if not rows and not parse_error:
        errors.append({"row": 0, "message": "文件没有可读取的数据行"})
    if headers and required - headers:
        errors.append({"row": 0, "message": "缺少字段：" + "、".join(sorted(required - headers))})
    if payload.data_kind == "inventory" and not errors:
        normalized, row_errors, duplicates = inventory_rows(rows)
        errors.extend(row_errors)
        if not errors:
            result = store.import_inventory(payload.filename, normalized, payload.as_of_date.isoformat() if payload.as_of_date else None, payload.actor_id, tenant, duplicates)
            result["summary"] = json.loads(result.pop("summary_json"))
            result["errors"] = json.loads(result.pop("errors_json"))
            return result
    status = "failed" if errors else "validated_only"
    summary = {"rows": len(rows), "fields": sorted(headers), "data_kind": payload.data_kind, "data_kind_label": schema["label"], "sheet_name": payload.sheet_name or None, "next_step": "仅完成格式校验；此类型尚未接入正式计算快照。"}
    result = store.add_data_import(payload.filename, payload.mode, status, summary, errors, payload.actor_id, tenant)
    result["summary"] = json.loads(result.pop("summary_json"))
    result["errors"] = json.loads(result.pop("errors_json"))
    return result


@app.get("/api/v1/cases")
@numeric_response
def list_cases(risk_type: Optional[str] = Query(default=None), x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    return {"items": store.cases(tenant_from_header(x_tenant_id), risk_type)}


@app.post("/api/v1/demo/reset")
@numeric_response
def reset_demo() -> Dict[str, Any]:
    if os.environ.get("INVENTORY_AGENT_MODE", "demo") != "demo" or os.environ.get("INVENTORY_AGENT_ALLOW_DEMO_RESET") != "1":
        raise HTTPException(403, "演示重置默认关闭，仅可在隔离演示环境显式启用")
    try:
        store.reset_demo()
    except ValueError as exc:
        _service_error(exc)
    return {"ok": True, "snapshot_id": "snapshot-demo-v1", "source_label": "合成样例回放"}


def _retail_demo(tenant: str) -> Optional[Dict[str, Any]]:
    # 真实快照优先于启动时的默认 demo 模式，不能给导入数据补上模拟账户。
    if store.real_inventory_snapshot(tenant) or os.environ.get("INVENTORY_AGENT_MODE", "demo") == "real" or tenant != "demo":
        return None
    sample = store.one("SELECT id FROM snapshots WHERE id='snapshot-demo-v1' AND tenant_id=? AND is_sample=1 AND source='sample-data/ac01'", (tenant,))
    return demo_dataset(store.risks(tenant), TRANSFER_NETWORK_STORES) if sample else None


def _retail_pending_approvals(tenant: str, snapshot_id: Optional[str], store_id: Optional[str]) -> Dict[str, Any]:
    """按待审批方案当前版本涉及的商品成本汇总，不把它解释为现金收益。"""
    total = Decimal(0)
    count = missing_count = 0
    proposal_rows = store.rows(
        "SELECT id FROM proposals WHERE tenant_id=? AND status='pending_approval' AND snapshot_id=?",
        (tenant, snapshot_id),
    ) if snapshot_id else []
    for row in proposal_rows:
        proposal = store.proposal(row["id"], tenant)
        risk = store.risk(proposal["risk_id"], tenant) if proposal else None
        if not proposal or (store_id and store_id != "all" and (not risk or risk.get("store_id") != store_id)):
            continue
        count += 1
        payload = ((proposal.get("version") or {}).get("payload") or {})
        input_data = payload.get("input") or {}
        calculation = payload.get("calculation") or {}
        action = (payload.get("actions") or [{}])[0]
        kind = payload.get("proposal_type")
        if kind == "expiry-rescue":
            forecast = calculation.get("forecast") or {}
            quantities = [forecast.get(f"{key}_qty", input_data.get(f"{key}_qty")) for key in ("transfer", "promo", "return")]
            quantity = sum(Decimal(str(value or 0)) for value in quantities) if all(value is not None for value in quantities) else None
        elif kind == "procurement-brake":
            quantity = input_data.get("adjustment_qty")
        else:
            quantity = action.get("quantity") or (calculation.get("allocation") or {}).get("quantity") or input_data.get("quantity")
        unit_cost = action.get("unit_cost") or input_data.get("unit_cost") or (risk or {}).get("unit_cost")
        try:
            if quantity is None or unit_cost is None:
                raise ValueError("方案数量或成本缺失")
            total += Decimal(str(quantity)) * Decimal(str(unit_cost))
        except (ValueError, TypeError, ArithmeticError):
            missing_count += 1
    return {"amount": round(float(total), 2) if missing_count == 0 else None,
            "known_amount": round(float(total), 2), "count": count, "missing_count": missing_count,
            "basis": "当前待审批方案涉及的商品成本；不是预计回款或确定支出"}


def _retail_real_overview(tenant: str, period: int, store_id: Optional[str]) -> Dict[str, Any]:
    snapshot = store.real_inventory_snapshot(tenant)
    all_rows = store.real_inventory_store_summary(snapshot["id"], tenant) if snapshot else []
    if store_id and store_id != "all" and store_id not in {row["org_code"] for row in all_rows}:
        raise HTTPException(status_code=422, detail="门店不存在或不在当前数据范围内")
    rows = [row for row in all_rows if not store_id or store_id == "all" or row["org_code"] == store_id]
    source = snapshot["source"] if snapshot else "未接入经营数据"
    metadata = (snapshot or {}).get("metadata") or {}
    baseline_ready = bool(snapshot) and not metadata.get("teacher_baseline_missing_fields")
    # 当源字段缺失时保持 null，不把真实库存快照当作销售或账户流水。
    missing = ["bank_account_ledger", "sales_revenue", "sales_returns", "period_sales_cost", "inventory_history"]
    if not baseline_ready:
        missing.append("complete_risk_inputs")
    attention_by_store: Dict[str, List[Dict[str, Any]]] = {}
    if baseline_ready:
        for line in store.real_inventory_attention_items(snapshot["id"], tenant):
            attention_by_store.setdefault(line["org_code"], []).append({
                "risk_id": line["id"], "sku": line["sku"], "product": line["product_name"],
                "inventory_cost": round(float(line.get("cost_amount") or 0), 2),
                "risk_label": "滞销待核查", "priority": line.get("teacher_priority"),
                "reason": line.get("teacher_trigger_reason") or "需核查库存与销量差异",
            })
    performance = [{"store_id": row["org_code"], "store_name": row["org_name"], "sales": None, "gross_profit": None,
                    "gross_margin_pct": None, "inventory_cost": row["inventory_value"],
                    "risk_cost": row["candidate_inventory_value"] if baseline_ready else None,
                    "turnover_days": None, "risk_count": row["slow_moving_skus"] if baseline_ready else None,
                    "primary_risk": "滞销待核查" if attention_by_store.get(row["org_code"]) else "暂无重点风险" if baseline_ready else None,
                    "risk_types": ["滞销"] if attention_by_store.get(row["org_code"]) else [],
                    "risk_items": attention_by_store.get(row["org_code"], [])} for row in rows]
    performance.sort(key=lambda row: (row["risk_cost"] is not None, row["risk_cost"] or 0, row["inventory_cost"] or 0), reverse=True)
    stockout_filter = " AND org_code=?" if store_id and store_id != "all" else ""
    snapshot_id = snapshot["id"] if snapshot else None
    stockout_params = (snapshot_id, tenant, store_id) if stockout_filter else (snapshot_id, tenant)
    stockout_risk_count = store.one(
        "SELECT COUNT(*) AS n FROM real_inventory_lines WHERE snapshot_id=? AND tenant_id=?"
        + stockout_filter + " AND (COALESCE(teacher_stockout,0)=1 OR COALESCE(teacher_near_stockout,0)=1)",
        stockout_params,
    )["n"] if baseline_ready else None
    return {
        "is_demo": False, "source": source, "as_of_date": (snapshot or {}).get("as_of_date"),
        "scope": {"store_id": store_id or "all", "store_count": len(rows), "period": period},
        "account": {"balance": None, "opening_balance": None, "cash_in": None, "cash_out": None,
                    "is_demo": False, "source": "未接入账户流水", "status": "unavailable"},
        "sales": {"amount": None, "gross_profit": None, "gross_margin_pct": None,
                  "amount_7d": None, "gross_margin_7d_pct": None, "change_pct": None},
        "inventory": {"cost": round(sum(row["inventory_value"] for row in rows), 2) if snapshot else None,
                      "risk_cost": round(sum(row["candidate_inventory_value"] for row in rows), 2) if baseline_ready else None,
                      "risk_count": sum(row["slow_moving_skus"] for row in rows) if baseline_ready else None,
                      "risk_store_count": sum(1 for row in rows if row["slow_moving_skus"] > 0) if baseline_ready else None,
                      "stockout_risk_count": stockout_risk_count,
                      "turnover_days": None,
                      "sku_count": (store.one("SELECT COUNT(DISTINCT sku) AS n FROM real_inventory_lines WHERE snapshot_id=? AND tenant_id=?" + (" AND org_code=?" if store_id and store_id != "all" else ""), (snapshot["id"], tenant, store_id) if store_id and store_id != "all" else (snapshot["id"], tenant))["n"] if snapshot else 0),
                      "store_sku_count": sum(row["categories"] for row in rows)},
        "purchase_commitments": {"amount": None, "count": None, "horizon_days": 30,
                                 "status": "unavailable", "source": "未接入已确认采购付款计划"},
        "pending_approvals": _retail_pending_approvals(tenant, snapshot["id"] if snapshot else None, store_id),
        "trend": [],
        "stores": performance,
        "metadata": {"is_demo": False, "source": source, "currency": "CNY", "missing": missing + ["confirmed_purchase_payment_schedule"],
                     "assumptions": ["库存只展示最新已导入快照；所选天数不改变库存时点。", "源文件未提供的销售收入、账户资金和趋势保持不可用。"]},
    }


@app.get("/api/v1/retail/overview")
@numeric_response
def retail_overview(period: int = Query(default=30), store_id: Optional[str] = Query(default="all"),
                    x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    if period not in (7, 30):
        raise HTTPException(status_code=422, detail="period 仅支持 7 或 30 天")
    tenant = tenant_from_header(x_tenant_id)
    dataset = _retail_demo(tenant)
    if not dataset:
        return _retail_real_overview(tenant, period, store_id)
    try:
        result = demo_overview(dataset, period, store_id)
        result["pending_approvals"] = _retail_pending_approvals(tenant, "snapshot-demo-v1", store_id)
        return result
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error))


@app.get("/api/v1/retail/simulation-options")
@numeric_response
def retail_simulation_options(x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    dataset = _retail_demo(tenant)
    if dataset:
        return simulation_options(dataset)
    snapshot = store.real_inventory_snapshot(tenant)
    rows = store.real_inventory_store_summary(snapshot["id"], tenant) if snapshot else []
    return {"stores": [{"id": row["org_code"], "name": row["org_name"]} for row in rows], "categories": [],
            "is_demo": False, "source": (snapshot or {}).get("source", "未接入经营数据"),
            "as_of_date": (snapshot or {}).get("as_of_date")}


@app.post("/api/v1/retail/simulate")
@numeric_response
def retail_simulate(payload: RetailSimulationInput, x_tenant_id: Optional[str] = Header(default=None)) -> Dict[str, Any]:
    tenant = tenant_from_header(x_tenant_id)
    dataset = _retail_demo(tenant)
    if not dataset:
        options = retail_simulation_options(x_tenant_id)
        if payload.store_id and payload.store_id != "all" and payload.store_id not in {row["id"] for row in options["stores"]}:
            raise HTTPException(status_code=422, detail="门店不存在或不在当前数据范围内")
        return {"status": "unavailable", "is_demo": False, "source": options["source"], "scenario": payload.model_dump(mode="json"),
                "metrics": None, "weekly": [], "risks": [], "lines": [],
                "metadata": {"is_demo": False, "source": options["source"], "units": {"money": "CNY", "risk_count": "store_sku"},
                             "missing": ["adjustable_purchase_orders", "confirmed_in_transit", "forecast_demand", "supplier_payment_schedule"],
                             "assumptions": ["真实库存快照不足以推算未来采购支出和缺货风险，需要补齐采购、在途、需求与付款计划。"]}}
    try:
        return simulate_purchase(dataset, payload.horizon_days, payload.reduction_pct, payload.store_id,
                                 payload.category, payload.request_text)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error))


# 保持现有静态 Demo 作为同一个产品入口；生产化时可替换为 Next.js build。
app.mount("/assets", StaticFiles(directory=str(ROOT / "assets")), name="assets")


@app.get("/", include_in_schema=False)
def frontend_index():
    return FileResponse(ROOT / "index.html")


@app.get("/{filename}", include_in_schema=False)
def frontend_file(filename: str):
    allowed = {"index.html", "app.js", "styles.css", "favicon.svg", "retail-app.js", "retail-app.css", "retail-workbenches.js", "retail-workbenches.css", "retail-simulation.js", "retail-simulation.css"}
    if filename not in allowed:
        raise HTTPException(404)
    return FileResponse(ROOT / filename)
