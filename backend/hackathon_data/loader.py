"""Validated, time-scoped retail fixtures; future outcomes have a separate reader.

The fixture directory is a trusted deployment input, never an HTTP upload path.
Every read model is a copy. Scenario changes therefore cannot mutate BASE or a
different session, and expected-answer workbooks are never opened at runtime.
"""
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
from functools import lru_cache
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_ROOT = Path(__file__).resolve().parents[2] / "sample-data" / "retail-v2"
BUSINESS_TIMEZONE = timezone(timedelta(hours=8))
SHEETS = {
    "门店": "stores", "商品": "products", "供应商": "suppliers", "人员": "people",
    "门店商品": "store_products", "风险规则": "risk_policies", "库存快照": "inventory",
    "商品批次": "lots", "逐日销售": "sales_daily", "逐日可售状态": "availability_daily",
    "库存变动": "inventory_movements", "历史采购": "historical_orders",
    "逐日需求假设": "demand_forecasts", "路线与运费": "routes", "收货日历": "store_calendar",
    "已确认采购": "purchase_orders", "采购在途": "in_transit", "账户快照": "accounts",
    "应付款": "payables", "退供条款": "return_terms", "阶段促销": "promotion_stages",
    "组合商品": "bundle_items", "假设说明": "assumptions", "场景目录": "scenarios",
    "独立风险输入": "risk_inputs", "隔离场景覆盖": "overrides", "场景批次": "scenario_lots",
    "原始材料": "materials", "采购意向": "purchase_intents", "抵款场景应付": "scenario_payables",
}
REPLAY_SHEETS = {
    "方案及任务": "tasks", "业务回执": "business_receipts", "渠道动作记录": "channel_actions",
    "批次销售": "sales", "现金流水": "cash_events", "核销分配": "cash_allocations",
    "退供确认": "return_confirmations",
}
SCOPES = {"BASE", "isolated_risk", "isolated_return", "isolated_purchase", "isolated_feedback"}
INTERNAL_TABLES = {"scenarios", "overrides", "scenario_lots", "scenario_payables"}
OVERRIDE_FIELDS = {
    "inventory": {"quantity", "lot_id", "unit_cost", "sellable_until"},
    "demand_forecasts": {"demand_base"}, "store_products": {"safety_stock_qty"},
    "store_calendar": {"is_open", "can_receive"},
    "procurement_policy": {"min_order_qty", "order_multiple"},
    "transfer_policy": {"available_alternative_transfer_qty"},
}


def parse_clock(value):
    try:
        result = datetime.fromisoformat(value)
    except (ValueError, TypeError) as exc:
        raise ValueError("场景时钟必须是带时区的 ISO 时间") from exc
    if result.tzinfo is None or result.utcoffset() is None:
        raise ValueError("场景时钟必须包含时区")
    return result.astimezone(BUSINESS_TIMEZONE)


def _value(value, spec, context):
    if value is None:
        if not spec["nullable"]:
            raise ValueError(f"{context} 缺少必填值")
        return None
    kind = spec["type"]
    if kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return date.fromisoformat(str(value)).isoformat()
    if kind == "boolean":
        if not isinstance(value, bool):
            raise ValueError(f"{context} 必须是布尔值")
    elif kind == "number":
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError(f"{context} 必须是有限数值")
    elif kind == "text":
        if not isinstance(value, str):
            raise ValueError(f"{context} 必须是文本")
    else:
        raise ValueError(f"{context} 的字段类型不受支持")
    return value


def _signature(root, replay):
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    roles = {"future_replay"} if replay else {"facts_or_assumptions", "isolated_scenarios"}
    entries = [item for item in manifest["files"] if item["role"] in roles]
    paths = [root / "manifest.json", root / "data_dictionary.json"]
    for entry in entries:
        filename = entry["file"]
        if Path(filename).name != filename or "/" in filename or "\\" in filename:
            raise ValueError("数据清单中的文件名必须是当前目录下的文件")
        path = root / "xlsx" / filename
        if path.resolve().parent != (root / "xlsx").resolve():
            raise ValueError("数据文件不在声明的数据目录中")
        paths.append(path)
    return tuple((str(path), path.stat().st_mtime_ns, path.stat().st_size) for path in paths)


@lru_cache(maxsize=4)
def _read_package(root_text, replay, signature):
    root = Path(root_text)
    manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    dictionary = json.loads((root / "data_dictionary.json").read_text(encoding="utf-8"))
    specs = {(item["file"], item["sheet"], item["field"]): item for item in dictionary}
    roles = {"future_replay"} if replay else {"facts_or_assumptions", "isolated_scenarios"}
    mapping = REPLAY_SHEETS if replay else SHEETS
    tables = {}
    selected = [item for item in manifest["files"] if item["role"] in roles]
    for entry in selected:
        filename = entry["file"]
        path = root / "xlsx" / filename
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest != entry["sha256"] or path.stat().st_size != entry["bytes"]:
            raise ValueError(f"{filename} 与数据清单校验值不一致")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if workbook.sheetnames != [item["name"] for item in entry["sheets"]]:
                raise ValueError(f"{filename} 的工作表与清单不一致")
            for sheet in entry["sheets"]:
                name = sheet["name"]
                if name not in mapping or mapping[name] in tables:
                    raise ValueError(f"未声明或重复的数据表：{name}")
                iterator = workbook[name].iter_rows()
                header = [cell.value for cell in next(iterator)]
                if header != sheet["columns"] or len(header) != len(set(header)):
                    raise ValueError(f"{filename}/{name} 字段与清单不一致")
                rows = []
                for index, cells in enumerate(iterator, 2):
                    if len(cells) != len(header):
                        raise ValueError(f"{filename}/{name}:{index} 列数不一致")
                    row = {}
                    for field, cell in zip(header, cells):
                        if cell.data_type == "f":
                            raise ValueError(f"{filename}/{name}:{index} 原始输入不能包含公式")
                        context = f"{filename}/{name}:{index}/{field}"
                        spec = specs.get((filename, name, field))
                        if spec is None:
                            raise ValueError(f"{context} 未在字段字典中声明")
                        row[field] = _value(cell.value, spec, context)
                    if "known_at" in row:
                        parse_clock(row["known_at"])
                    rows.append(row)
                if len(rows) != sheet["row_count"]:
                    raise ValueError(f"{filename}/{name} 行数与清单不一致")
                tables[mapping[name]] = rows
        finally:
            workbook.close()
    if set(tables) != set(mapping.values()):
        raise ValueError("数据包缺少必需工作表")
    if not replay:
        _validate_references(tables)
    return manifest, tables


def _package(root=None, replay=False):
    path = Path(root or DEFAULT_ROOT).resolve()
    return _read_package(str(path), replay, _signature(path, replay))


def _unique(rows, fields, table):
    result = {}
    for row in rows:
        key = tuple(row[field] for field in fields)
        if any(value is None for value in key) or key in result:
            raise ValueError(f"{table} 标识为空或重复：{key}")
        result[key] = row
    return result


def _validate_references(tables):
    stores = _unique(tables["stores"], ("store_id",), "stores")
    products = _unique(tables["products"], ("sku_id",), "products")
    suppliers = _unique(tables["suppliers"], ("supplier_id",), "suppliers")
    lots = _unique(tables["lots"] + tables["scenario_lots"], ("lot_id",), "lots")
    for name in ("inventory", "sales_daily", "availability_daily", "store_products", "demand_forecasts",
                 "inventory_movements", "historical_orders", "in_transit", "purchase_orders"):
        for row in tables[name]:
            if (row["store_id"],) not in stores or (row["sku_id"],) not in products:
                raise ValueError(f"{name} 引用未知门店或商品")
            if row.get("lot_id"):
                lot = lots.get((row["lot_id"],))
                if not lot or lot["sku_id"] != row["sku_id"]:
                    raise ValueError(f"{name} 批次与商品不匹配")
    for name in ("products", "lots", "scenario_lots"):
        for row in tables[name]:
            if (row["supplier_id"],) not in suppliers:
                raise ValueError(f"{name} 引用未知供应商")
    unique_fields = {
        "inventory": ("inventory_id",), "sales_daily": ("sales_day_id",),
        "availability_daily": ("store_id", "sku_id", "date"),
        "store_products": ("store_id", "sku_id"), "in_transit": ("shipment_line_id",),
        "demand_forecasts": ("store_id", "sku_id", "date"), "routes": ("route_id",),
        "store_calendar": ("store_id", "date"), "scenarios": ("scenario_id",),
        "risk_inputs": ("scenario_id", "branch_id"), "overrides": ("override_id",),
    }
    for name, fields in unique_fields.items():
        _unique(tables[name], fields, name)
    _unique(tables["inventory"], ("store_id", "sku_id", "lot_id", "stock_state"), "inventory assets")
    shipments = {row["shipment_line_id"]: row for row in tables["in_transit"]}
    for row in tables["inventory"]:
        if row["quantity"] < 0 or min(row["blocked_qty"], row["reserved_qty"]) < 0:
            raise ValueError("库存数量不能为负数")
        if row["blocked_qty"] + row["reserved_qty"] > row["quantity"]:
            raise ValueError("冻结与预留合计超过实物库存")
        if row["stock_state"] == "in_transit":
            shipment = shipments.get(row["shipment_line_id"])
            if not shipment or any(row[key] != shipment[key] for key in ("store_id", "sku_id", "lot_id")):
                raise ValueError("在途库存缺少一致的运输明细")
            if row["quantity"] != shipment["shipped_qty"] - shipment["received_qty"]:
                raise ValueError("在途库存与运输未收数量不一致")
        elif row["stock_state"] != "on_hand":
            raise ValueError("未知库存状态")
    for row in tables["routes"]:
        if any((row[field],) not in stores for field in ("from_store_id", "to_store_id")):
            raise ValueError("路线引用未知门店")
    for row in tables["scenarios"]:
        if row["input_scope"] not in SCOPES:
            raise ValueError("场景读取范围未定义")


def _branches(scenario, tables):
    declared = scenario["branches"].split("|")
    risk = [row["branch_id"] for row in tables["risk_inputs"] if row["scenario_id"] == scenario["scenario_id"]]
    aliases = {}
    if scenario["input_scope"] == "isolated_risk" and len(risk) == 1:
        aliases = {branch: risk[0] for branch in declared if branch not in risk}
    choices = list(dict.fromkeys(aliases.get(branch, branch) for branch in declared + risk))
    return choices, aliases


def list_scenarios(*, root=None):
    manifest, tables = _package(root)
    items = []
    for row in tables["scenarios"]:
        choices, aliases = _branches(row, tables)
        items.append({key: row[key] for key in ("scenario_id", "title", "store_id", "sku_id", "lot_id", "decision_at", "evaluation_end")} |
                     {"branches": choices, "branch_aliases": aliases, "input_scope": row["input_scope"]})
    return {"dataset_id": manifest["dataset_id"], "decision_at": manifest["decision_at"], "is_demo": True,
            "complete_store_count": manifest["complete_store_count"],
            "directory_store_count": manifest["directory_store_count"], "items": items}


def _select(scenario_id, branch_id, root):
    manifest, tables = _package(root)
    scenario = next((row for row in tables["scenarios"] if row["scenario_id"] == scenario_id), None)
    if scenario is None:
        raise ValueError("场景不存在")
    choices, aliases = _branches(scenario, tables)
    requested = branch_id or choices[0]
    canonical = aliases.get(requested, requested)
    if canonical not in choices:
        raise ValueError("分支不属于所选场景")
    return manifest, tables, scenario, canonical, aliases


def _apply_override(tables, override):
    name, field = override["table_name"], override["field"]
    if field not in OVERRIDE_FIELDS.get(name, set()):
        raise ValueError("场景覆盖字段不在允许列表中")
    key = override["record_key"].split(":")
    if len(key) not in (2, 3):
        raise ValueError("场景覆盖的记录键不合法")
    value = json.loads(override["value_json"])
    if field in ("is_open", "can_receive"):
        if not isinstance(value, bool):
            raise ValueError("营业与收货状态覆盖必须是布尔值")
    elif field == "sellable_until":
        value = date.fromisoformat(value).isoformat()
    elif field == "lot_id":
        if not isinstance(value, str) or not any(row["lot_id"] == value for row in tables["lots"]):
            raise ValueError("覆盖批次尚未可知或未声明")
    elif isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        raise ValueError("数量与金额覆盖必须为非负有限数值")
    if name in ("procurement_policy", "transfer_policy"):
        rows = [row for row in tables[name] if row["store_id"] == key[0] and row["sku_id"] == key[1]]
        if not rows:
            rows = [{"store_id": key[0], "sku_id": key[1]}]
            tables[name].extend(rows)
    else:
        rows = [row for row in tables[name] if row.get("store_id") == key[0]]
        if name == "store_calendar":
            start, end = key[1].split("..")
            rows = [row for row in rows if start <= row["date"] <= end]
        else:
            rows = [row for row in rows if row.get("sku_id") == key[1]]
            if name == "inventory":
                rows = [row for row in rows if row["stock_state"] == "on_hand"]
                if len(rows) != 1:
                    raise ValueError("场景库存覆盖必须明确对应唯一在库批次")
            if len(key) == 3:
                start, end = key[2].split("..")
                rows = [row for row in rows if start <= row["date"] <= end]
    if not rows:
        raise ValueError("场景覆盖未找到已知记录")
    for row in rows:
        row[field] = value
        if name == "demand_forecasts" and field == "demand_base":
            # This fixture replaces only its base assumption. BASE's low/high
            # estimates no longer describe it and must not masquerade as bounds.
            row["demand_low"] = None
            row["demand_high"] = None
        row.setdefault("applied_override_ids", []).append(override["override_id"])


def load_scenario(scenario_id, branch_id=None, *, clock_at=None, confirmed_override_ids=(), root=None):
    manifest, raw, scenario, branch, aliases = _select(scenario_id, branch_id, root)
    clock = parse_clock(clock_at or scenario["decision_at"])
    if clock < parse_clock(scenario["known_at"]):
        raise ValueError("场景在当前时钟尚不可知")
    visible = lambda row: parse_clock(row["known_at"]) <= clock
    tables = {name: [deepcopy(row) for row in rows if visible(row) and
                    (not row.get("scenario_id") or row["scenario_id"] == scenario_id) and
                    (not row.get("branch_id") or row["branch_id"] == branch)]
              for name, rows in raw.items() if name not in INTERNAL_TABLES}
    tables["procurement_policy"], tables["transfer_policy"] = [], []
    tables["lots"].extend(deepcopy(row) for row in raw["scenario_lots"]
                          if row["scenario_id"] == scenario_id and visible(row))
    tables["payables"].extend(deepcopy(row) for row in raw["scenario_payables"]
                              if row["scenario_id"] == scenario_id and row["branch_id"] == branch and visible(row))
    confirmed = set(confirmed_override_ids)
    candidates = [row for row in raw["overrides"] if row["scenario_id"] == scenario_id and row["branch_id"] in (branch, "*")]
    if confirmed - {row["override_id"] for row in candidates if visible(row)}:
        raise ValueError("不能确认其他场景或当前时钟尚未知晓的覆盖项")
    applied, pending = [], []
    for row in candidates:
        if not visible(row):
            continue
        if scenario["input_scope"] == "isolated_feedback" and row["override_id"] not in confirmed:
            pending.append({key: row[key] for key in ("override_id", "table_name", "record_key", "field", "known_at")})
            continue
        _apply_override(tables, row)
        applied.append(row["override_id"])
    target = {key: scenario[key] for key in ("store_id", "sku_id", "lot_id")}
    isolated = tables["risk_inputs"]
    if isolated:
        # These are complete, independent risk inputs, including deliberate nulls.
        item = isolated[0]
        target["sku_id"] = item["sku_id"]
        tables["inventory"] = [{"inventory_id": f"{scenario_id}:{branch}", "store_id": target["store_id"],
                                "sku_id": item["sku_id"], "lot_id": target["lot_id"], "stock_state": "on_hand",
                                "quantity": item["quantity"], "blocked_qty": 0, "reserved_qty": 0,
                                "unit_cost": item["unit_cost"], "sellable_until": item["sellable_until"],
                                "shipment_line_id": None, "source": item["source"], "known_at": item["known_at"]}]
        tables["lots"] = [{"lot_id": target["lot_id"], "sku_id": item["sku_id"], "supplier_id": None,
                           "production_date": None, "expiry_date": None, "source": item["source"],
                           "known_at": item["known_at"], "input_role": "independent_risk_input"}]
        for name in ("sales_daily", "availability_daily", "inventory_movements", "historical_orders", "in_transit", "demand_forecasts"):
            tables[name] = []
    elif scenario["input_scope"] == "isolated_risk":
        raise ValueError("独立风险场景缺少自身输入，不能使用 BASE 补齐")
    references = {row.get(key) for name in ("return_terms", "purchase_orders", "purchase_intents", "payables", "lots")
                  for row in tables[name] for key in ("source_ref", "original_document_ref", "po_line_id")}
    tables["materials"] = [row for row in tables["materials"] if row["material_id"] in references or
                           (isolated and scenario["input_scope"] == "isolated_feedback" and
                            row["type"] == "manager_feedback" and row["party_id"] == target["store_id"])]
    warnings = [f"场景目录分支 {old} 映射到独立输入分支 {new}" for old, new in aliases.items()]
    if scenario["input_scope"] in ("isolated_return", "isolated_purchase"):
        warnings.append("本场景使用独立期初库存，BASE 历史余额不能视为该快照的对账结果")
    identity = {"dataset_id": manifest["dataset_id"], "scenario_id": scenario_id, "branch_id": branch,
                "clock_at": clock.isoformat(), "applied_override_ids": sorted(applied),
                "files": [entry["sha256"] for entry in manifest["files"] if entry["role"] in {"facts_or_assumptions", "isolated_scenarios"}]}
    snapshot = "SCENARIO-" + hashlib.sha256(json.dumps(identity, sort_keys=True).encode()).hexdigest()[:24]
    return {"dataset_id": manifest["dataset_id"], "scenario_id": scenario_id, "branch_id": branch,
            "canonical_branch_id": branch, "clock_at": clock.isoformat(), "snapshot_id": snapshot,
            "data_version": snapshot, "is_demo": True, "currency": "CNY", "amount_unit": "yuan",
            "target": target, "evaluation_end": scenario["evaluation_end"], "input_scope": scenario["input_scope"],
            "tables": tables, "warnings": warnings, "applied_override_ids": applied,
            "pending_overrides": pending}


def load_replay(scenario_id, branch_id=None, *, clock_at, root=None):
    """Execution-only fixtures. The caller must check approval and compatibility."""
    _, _, _, branch, _ = _select(scenario_id, branch_id, root)
    clock = parse_clock(clock_at)
    _, raw = _package(root, replay=True)
    tables = {name: [deepcopy(row) for row in rows if row["scenario_id"] == scenario_id and
                    row["branch_id"] == branch and parse_clock(row["known_at"]) <= clock and
                    (not row.get("occurred_at") or parse_clock(row["occurred_at"]) <= clock)]
              for name, rows in raw.items()}
    return {"scenario_id": scenario_id, "branch_id": branch, "clock_at": clock.isoformat(), "is_demo": True, "tables": tables}
